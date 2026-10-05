"""List companies shared by Trucost, LSEG, and FactSet for one year."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from rapidfuzz import fuzz, process


PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from src.dataset_readers.cdp_read import find_header_row, get_col_idx, parse_year_file
from src.dataset_readers.common_companies_by_year import (
    CDP_RAW_ROOT,
    find_common_companies,
    normalize_isin,
    normalize_name,
)


CDP_NAME_THRESHOLD = 98.0
ISIN_PATTERN = re.compile(r"\b[A-Z]{2}[A-Z0-9]{9}[0-9]\b")


def extract_isins(value: object) -> set[str]:
    if pd.isna(value):
        return set()
    return {
        normalize_isin(match)
        for match in ISIN_PATTERN.findall(str(value).upper())
    }


def read_legacy_cdp_organizations(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    summary_name = (
        "Summary Data"
        if "Summary Data" in workbook.sheetnames
        else "Summary"
        if "Summary" in workbook.sheetnames
        else None
    )
    if summary_name is None:
        workbook.close()
        return pd.DataFrame(columns=["cdp_name", "cdp_isins"])

    sheet = workbook[summary_name]
    header_row = find_header_row(sheet)
    headers = [
        sheet.cell(row=header_row, column=column).value
        for column in range(1, sheet.max_column + 1)
    ]
    headers = ["" if value is None else str(value) for value in headers]
    name_index = get_col_idx(
        headers,
        ("organization", "organisation", "response organisation"),
    )
    isin_index = get_col_idx(headers, ("primary isin", "isins", "isin"))
    rows = []
    if name_index is not None:
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            name = row[name_index] if name_index < len(row) else None
            if pd.isna(name) or not str(name).strip():
                continue
            isin_value = (
                row[isin_index]
                if isin_index is not None and isin_index < len(row)
                else None
            )
            rows.append(
                {
                    "cdp_name": str(name).strip(),
                    "cdp_isins": extract_isins(isin_value),
                }
            )
    workbook.close()
    return pd.DataFrame(rows).drop_duplicates("cdp_name")


def read_2024_cdp_organizations(year_dir: Path) -> pd.DataFrame:
    summary_files = sorted(year_dir.rglob("*summary*.parquet"))
    if not summary_files:
        return pd.DataFrame(columns=["cdp_name", "cdp_isins"])
    source = pd.read_parquet(summary_files[0])
    lookup = {str(column).casefold(): column for column in source.columns}
    name_column = next(
        (
            lookup[candidate]
            for candidate in [
                "disclosing_organization",
                "organization_name",
                "organization",
            ]
            if candidate in lookup
        ),
        None,
    )
    isin_column = next(
        (
            lookup[candidate]
            for candidate in ["isin", "primary_isin", "isins"]
            if candidate in lookup
        ),
        None,
    )
    if name_column is None:
        return pd.DataFrame(columns=["cdp_name", "cdp_isins"])
    output = pd.DataFrame(
        {
            "cdp_name": source[name_column].astype("string").str.strip(),
            "cdp_isins": (
                source[isin_column].map(extract_isins)
                if isin_column is not None
                else [set() for _ in range(len(source))]
            ),
        }
    )
    return output.dropna(subset=["cdp_name"]).drop_duplicates("cdp_name")


def read_cdp_organizations(year: int, cdp_root: Path) -> pd.DataFrame:
    year_dir = cdp_root / str(year)
    if year == 2024:
        organizations = read_2024_cdp_organizations(year_dir)
        if not organizations.empty:
            return organizations
    path = parse_year_file(year_dir)
    if path is None or path.suffix.casefold() == ".parquet":
        return pd.DataFrame(columns=["cdp_name", "cdp_isins"])
    return read_legacy_cdp_organizations(path)


def add_cdp_names(
    companies: pd.DataFrame, cdp_organizations: pd.DataFrame
) -> pd.DataFrame:
    output = companies.copy()
    output["in_cdp"] = False
    output["cdp_name"] = pd.NA
    output["cdp_match_method"] = pd.NA
    output["cdp_name_score"] = pd.NA
    if cdp_organizations.empty:
        return output

    isin_to_name = {
        isin: row.cdp_name
        for row in cdp_organizations.itertuples(index=False)
        for isin in row.cdp_isins
    }
    names_by_key = (
        cdp_organizations.assign(
            cdp_name_key=cdp_organizations["cdp_name"].map(normalize_name)
        )
        .drop_duplicates("cdp_name_key")
        .set_index("cdp_name_key")["cdp_name"]
        .to_dict()
    )
    choices = list(names_by_key)

    for index, row in output.iterrows():
        company_isin = normalize_isin(row.get("isin"))
        if company_isin in isin_to_name:
            output.at[index, "in_cdp"] = True
            output.at[index, "cdp_name"] = isin_to_name[company_isin]
            output.at[index, "cdp_match_method"] = "year+isin"
            output.at[index, "cdp_name_score"] = 100.0
            continue

        aliases = {
            normalize_name(row.get("trucost_name")),
            normalize_name(row.get("lseg_name")),
            normalize_name(row.get("factset_name")),
        } - {""}
        best_match = None
        for alias in aliases:
            match = process.extractOne(
                alias,
                choices,
                scorer=fuzz.WRatio,
                score_cutoff=CDP_NAME_THRESHOLD,
            )
            if match is not None and (
                best_match is None or match[1] > best_match[1]
            ):
                best_match = match
        if best_match is None:
            continue
        name_key, score, _ = best_match
        output.at[index, "in_cdp"] = True
        output.at[index, "cdp_name"] = names_by_key[name_key]
        output.at[index, "cdp_match_method"] = "year+fuzzy_name"
        output.at[index, "cdp_name_score"] = float(score)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "List companies shared by Trucost, LSEG, and FactSet and show "
            "their CDP name when present."
        )
    )
    parser.add_argument("year", type=int, help="Fiscal/reporting year.")
    parser.add_argument(
        "--cdp-root",
        type=Path,
        default=CDP_RAW_ROOT,
        help="Root directory containing CDP year folders.",
    )
    parser.add_argument("--output", type=Path, help="Output CSV path.")
    args = parser.parse_args()

    companies = find_common_companies(
        years=[args.year],
        cdp_root=None,
        names_only=True,
    )
    result = add_cdp_names(
        companies,
        read_cdp_organizations(args.year, args.cdp_root),
    )
    result = result.sort_values(["trucost_name", "isin"], na_position="last")
    output_path = args.output or (
        PROJECT_ROOT
        / f"data/processed/common_company_names_{args.year}.csv"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    result.to_csv(output_path, index=False)
    #print(result.to_string(index=False))
    print(f"\nSaved {len(result):,} companies to {output_path}")


if __name__ == "__main__":
    main()
