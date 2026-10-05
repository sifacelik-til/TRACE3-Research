from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Dict, List

import pandas as pd
from openpyxl import load_workbook
from rapidfuzz import fuzz, process

from src.cdp_extraction.extract_cdp_org_answers_2014_2024 import (
    DEFAULT_CDP_ROOT,
    find_header_row,
    get_col_idx,
    is_missing,
    parse_year_file,
)

FACTSET_RAW_ROOT = Path(r"/data/raw/FactSet")
TRUCOST_RAW_ROOT = Path(r"/data/raw/Trucost (Access through WRDS)")
LSEG_RAW_ROOT = Path(r"/data/raw/LSEG")


TOKEN_MAP = {
    "co": "company",
    "company": "company",
    "corp": "corporation",
    "corporation": "corporation",
    "ltd": "limited",
    "limited": "limited",
}


def normalize_name(value: object) -> str:
    text = "" if pd.isna(value) else str(value).casefold()
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    tokens = [token for token in text.split() if token]
    tokens = [TOKEN_MAP.get(token, token) for token in tokens]
    return " ".join(tokens)


def _clean_year(value: object) -> int | None:
    parsed = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    if pd.isna(parsed):
        return None
    year = int(parsed)
    if year < 1900 or year > 2100:
        return None
    return year


def _clean_isin(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    if not text or text.lower() in {"nan", "none", "null"}:
        return ""
    return text


def _clean_company_id(value: object) -> str:
    parsed = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    if pd.isna(parsed):
        text = str(value).strip()
        return "" if text.lower() in {"", "nan", "none", "null"} else text
    return str(int(parsed))


def _first_non_empty(series: pd.Series) -> str:
    for value in series:
        if value is None:
            continue
        text = str(value).strip()
        if text and text.lower() not in {"nan", "none"}:
            return text
    return ""


def _read_pipe(path: Path, usecols: List[str]) -> pd.DataFrame:
    try:
        return pd.read_csv(path, sep="|", dtype=str, usecols=usecols, low_memory=False, encoding="utf-8")
    except UnicodeDecodeError:
        return pd.read_csv(path, sep="|", dtype=str, usecols=usecols, low_memory=False, encoding="latin-1")


def build_reference_from_current(df: pd.DataFrame) -> pd.DataFrame:
    required = {"factset_entity_id", "factset_company_name", "factset_ticker"}
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"Missing required FactSet columns: {sorted(missing)}")

    ref = df[["factset_entity_id", "factset_company_name", "factset_ticker", "trucost_isin"]].fillna("").copy()
    ref["factset_entity_id"] = ref["factset_entity_id"].astype(str).str.strip()
    ref["factset_company_name"] = ref["factset_company_name"].astype(str).str.strip()
    ref["factset_ticker"] = ref["factset_ticker"].astype(str).str.strip()
    ref["trucost_isin"] = ref["trucost_isin"].astype(str).str.strip().str.upper()
    ref = ref[(ref["factset_entity_id"] != "") & (ref["factset_company_name"] != "")]
    ref["name_norm"] = ref["factset_company_name"].map(normalize_name)
    return ref.drop_duplicates()


def choose_candidate(candidates: pd.DataFrame, trucost_norm: str) -> tuple[str, str, str, float, str]:
    exact = candidates[candidates["name_norm"] == trucost_norm]
    if not exact.empty:
        row = exact.iloc[0]
        return row["factset_entity_id"], row["factset_company_name"], row["factset_ticker"], 1000.0, "exact"

    scores = candidates["name_norm"].map(lambda n: float(fuzz.token_sort_ratio(trucost_norm, n)))
    best_idx = scores.idxmax()
    row = candidates.loc[best_idx]
    return row["factset_entity_id"], row["factset_company_name"], row["factset_ticker"], float(scores.loc[best_idx]), "fuzzy"


def rematch_trucost_factset(df: pd.DataFrame) -> pd.DataFrame:
    required = {"trucost_companyname", "trucost_isin", "factset_entity_id", "factset_company_name", "factset_ticker"}
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"Missing required columns for rematch: {sorted(missing)}")

    ref = build_reference_from_current(df)
    isin_groups = {k: g for k, g in ref[ref["trucost_isin"] != ""].groupby("trucost_isin", sort=False)}
    name_groups = {k: g for k, g in ref.groupby("name_norm", sort=False)}
    all_norm_names = list(name_groups.keys())

    pairs = df[["trucost_companyname", "trucost_isin"]].fillna("").drop_duplicates().copy()
    pairs["trucost_isin"] = pairs["trucost_isin"].astype(str).str.strip().str.upper()
    pairs["trucost_norm"] = pairs["trucost_companyname"].map(normalize_name)

    mapped: List[Dict[str, object]] = []
    total = len(pairs)
    for i, (_, row) in enumerate(pairs.iterrows(), start=1):
        t_name = row["trucost_companyname"]
        t_isin = row["trucost_isin"]
        t_norm = row["trucost_norm"]

        if t_isin and t_isin in isin_groups:
            ent, name, tick, score, kind = choose_candidate(isin_groups[t_isin], t_norm)
            mtype = "isin+name" if kind == "exact" else "isin"
        else:
            exact_group = name_groups.get(t_norm)
            if exact_group is not None and not exact_group.empty:
                ent, name, tick, score, _ = choose_candidate(exact_group, t_norm)
                mtype = "exact"
            elif all_norm_names:
                best = process.extractOne(t_norm, all_norm_names, scorer=fuzz.token_sort_ratio)
                if best is None:
                    ent, name, tick, score, mtype = "", "", "", 0.0, "unmatched"
                else:
                    ent, name, tick, score, _ = choose_candidate(name_groups[best[0]], t_norm)
                    mtype = "fuzzy"
            else:
                ent, name, tick, score, mtype = "", "", "", 0.0, "unmatched"

        mapped.append(
            {
                "trucost_companyname": t_name,
                "trucost_isin": t_isin,
                "factset_entity_id_new": ent,
                "factset_company_name_new": name,
                "factset_ticker_new": tick,
                "fuzzy_score_new": score,
                "match_type_new": mtype,
                "source_name_used_new": t_name,
            }
        )
        if i % 2000 == 0 or i == total:
            print(f"  Matched {i}/{total} unique Trucost company/ISIN pairs", flush=True)

    map_df = pd.DataFrame(mapped)
    out = df.copy()
    out["trucost_isin"] = out["trucost_isin"].fillna("").astype(str).str.strip().str.upper()
    out = out.merge(map_df, on=["trucost_companyname", "trucost_isin"], how="left")

    out["factset_entity_id"] = out["factset_entity_id_new"].fillna("")
    out["factset_company_name"] = out["factset_company_name_new"].fillna("")
    out["factset_ticker"] = out["factset_ticker_new"].fillna("")
    out["fuzzy_score"] = pd.to_numeric(out["fuzzy_score_new"], errors="coerce").fillna(0.0)
    out["match_type"] = out["match_type_new"].fillna("unmatched")
    out["source_name_used"] = out["source_name_used_new"].fillna(out["trucost_companyname"])
    return out.drop(columns=[c for c in ["factset_entity_id_new", "factset_company_name_new", "factset_ticker_new", "fuzzy_score_new", "match_type_new", "source_name_used_new"] if c in out.columns])


def update_cdp_match_scores(df: pd.DataFrame) -> pd.DataFrame:
    required = {"cdp_org_name", "factset_company_name"}
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"Missing required columns for cdp_match_score update: {sorted(missing)}")

    out = df.copy()
    left = out["cdp_org_name"].map(normalize_name)
    right = out["factset_company_name"].map(normalize_name)
    exact_equivalent = left.eq(right)

    if "cdp_match_score" not in out.columns:
        out["cdp_match_score"] = pd.NA

    fuzzy = pd.to_numeric(out.get("fuzzy_score"), errors="coerce")
    old_score = pd.to_numeric(out["cdp_match_score"], errors="coerce")
    fallback = fuzzy.where(fuzzy.notna(), old_score.where(old_score.notna(), 0)).clip(lower=0, upper=999)
    out.loc[exact_equivalent, "cdp_match_score"] = 1000.0
    out.loc[~exact_equivalent, "cdp_match_score"] = fallback.loc[~exact_equivalent]
    return out


def _read_cdp_unique_orgs_from_xlsx(file_path: Path, year: int) -> List[Dict[str, object]]:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    summary_sheet_name = "Summary Data" if "Summary Data" in wb.sheetnames else ("Summary" if "Summary" in wb.sheetnames else None)
    if summary_sheet_name is None:
        return []

    ws = wb[summary_sheet_name]
    header_row = find_header_row(ws)
    headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
    headers = ["" if h is None else str(h) for h in headers]
    org_idx = get_col_idx(headers, ("organization", "account_name", "organisation", "response organisation", "disclosing organization", "discloser"))
    if org_idx is None:
        return []
    sector_idx = get_col_idx(headers, ("primary sector", "sectors", "questionnaire sector"))
    industry_idx = get_col_idx(headers, ("primary industry", "industry"))

    rows: List[Dict[str, object]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if org_idx >= len(row):
            continue
        org = row[org_idx]
        if is_missing(org):
            continue
        name = str(org).strip()
        sector = ""
        industry = ""
        if sector_idx is not None and sector_idx < len(row) and not is_missing(row[sector_idx]):
            sector = str(row[sector_idx]).strip()
        if industry_idx is not None and industry_idx < len(row) and not is_missing(row[industry_idx]):
            industry = str(row[industry_idx]).strip()

        rows.append(
            {
                "year": year,
                "source": "cdp",
                "company_name": name,
                "company_name_norm": normalize_name(name),
                "sector_description": sector,
                "industry_description": industry,
            }
        )

    if not rows:
        return []
    out = pd.DataFrame(rows).drop_duplicates(subset=["year", "source", "company_name_norm", "sector_description", "industry_description"])
    return out.to_dict("records")


def load_cdp_unique_names_by_year(cdp_root: Path, years: List[int]) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    for year in years:
        year_dir = cdp_root / str(year)
        file_path = parse_year_file(year_dir)
        if file_path is None:
            continue

        if file_path.suffix.lower() == ".parquet":
            cols = ["disclosing_organization", "primary_industry_name"]
            df = pd.read_parquet(file_path, columns=cols)
            df = df.rename(columns={"disclosing_organization": "company_name", "primary_industry_name": "industry_description"})
            df["company_name"] = df["company_name"].fillna("").astype(str).str.strip()
            df["industry_description"] = df["industry_description"].fillna("").astype(str).str.strip()
            df = df[df["company_name"] != ""].copy()
            df["year"] = year
            df["source"] = "cdp"
            df["company_name_norm"] = df["company_name"].map(normalize_name)
            df["sector_description"] = ""
            rows.extend(df[["year", "source", "company_name", "company_name_norm", "sector_description", "industry_description"]].drop_duplicates().to_dict("records"))
        else:
            rows.extend(_read_cdp_unique_orgs_from_xlsx(file_path, year))
    return pd.DataFrame(rows)


def load_factset_entity_descriptions(factset_root: Path) -> pd.DataFrame:
    entity_sector = _read_pipe(
        factset_root / "sym_entity_v1_full_12328" / "sym_entity_sector.txt",
        ["FACTSET_ENTITY_ID", "INDUSTRY_CODE", "SECTOR_CODE"],
    )
    sector_map = _read_pipe(
        factset_root / "ref_hub_v2_full_3565" / "factset_sector_map.txt",
        ["FACTSET_SECTOR_CODE", "FACTSET_SECTOR_DESC"],
    )
    industry_map = _read_pipe(
        factset_root / "ref_hub_v2_full_3565" / "factset_industry_map.txt",
        ["FACTSET_INDUSTRY_CODE", "FACTSET_INDUSTRY_DESC", "FACTSET_SECTOR_CODE"],
    )

    entity_sector["FACTSET_ENTITY_ID"] = entity_sector["FACTSET_ENTITY_ID"].fillna("").astype(str).str.strip()
    entity_sector["INDUSTRY_CODE"] = entity_sector["INDUSTRY_CODE"].fillna("").astype(str).str.strip()
    entity_sector["SECTOR_CODE"] = entity_sector["SECTOR_CODE"].fillna("").astype(str).str.strip()
    sector_map["FACTSET_SECTOR_CODE"] = sector_map["FACTSET_SECTOR_CODE"].fillna("").astype(str).str.strip()
    sector_map["FACTSET_SECTOR_DESC"] = sector_map["FACTSET_SECTOR_DESC"].fillna("").astype(str).str.strip()
    industry_map["FACTSET_INDUSTRY_CODE"] = industry_map["FACTSET_INDUSTRY_CODE"].fillna("").astype(str).str.strip()
    industry_map["FACTSET_INDUSTRY_DESC"] = industry_map["FACTSET_INDUSTRY_DESC"].fillna("").astype(str).str.strip()
    industry_map["FACTSET_SECTOR_CODE"] = industry_map["FACTSET_SECTOR_CODE"].fillna("").astype(str).str.strip()

    merged = entity_sector.merge(
        industry_map[["FACTSET_INDUSTRY_CODE", "FACTSET_INDUSTRY_DESC", "FACTSET_SECTOR_CODE"]],
        left_on="INDUSTRY_CODE",
        right_on="FACTSET_INDUSTRY_CODE",
        how="left",
    )
    merged["SECTOR_CODE_FINAL"] = merged["SECTOR_CODE"]
    merged.loc[merged["SECTOR_CODE_FINAL"] == "", "SECTOR_CODE_FINAL"] = merged["FACTSET_SECTOR_CODE"]
    merged = merged.merge(
        sector_map[["FACTSET_SECTOR_CODE", "FACTSET_SECTOR_DESC"]],
        left_on="SECTOR_CODE_FINAL",
        right_on="FACTSET_SECTOR_CODE",
        how="left",
    )

    out = merged[["FACTSET_ENTITY_ID", "FACTSET_SECTOR_DESC", "FACTSET_INDUSTRY_DESC"]].copy()
    out = out.rename(columns={"FACTSET_SECTOR_DESC": "sector_description", "FACTSET_INDUSTRY_DESC": "industry_description"})
    out["sector_description"] = out["sector_description"].fillna("").astype(str).str.strip()
    out["industry_description"] = out["industry_description"].fillna("").astype(str).str.strip()
    return out.drop_duplicates(subset=["FACTSET_ENTITY_ID"])


def load_trucost_descriptions(trucost_root: Path) -> pd.DataFrame:
    files = sorted(trucost_root.glob("*trucost*.csv"))
    if not files:
        return pd.DataFrame(columns=["year", "company_name_norm", "sector_description", "industry_description"])

    parts: List[pd.DataFrame] = []
    for path in files:
        try:
            df = pd.read_csv(path, usecols=["companyname", "fiscalyear", "simpleindustry", "tcprimarysectorid"], low_memory=False)
        except ValueError:
            try:
                df = pd.read_csv(path, usecols=["companyname", "fiscalyear", "simpleindustry"], low_memory=False)
                df["tcprimarysectorid"] = ""
            except ValueError:
                continue
        df["year"] = df["fiscalyear"].map(_clean_year)
        df["company_name"] = df["companyname"].fillna("").astype(str).str.strip()
        df["company_name_norm"] = df["company_name"].map(normalize_name)
        df["sector_description"] = df["tcprimarysectorid"].fillna("").astype(str).str.strip()
        df["industry_description"] = df["simpleindustry"].fillna("").astype(str).str.strip()
        df = df[(df["year"].notna()) & (df["company_name"] != "")]
        parts.append(df[["year", "company_name_norm", "sector_description", "industry_description"]])

    if not parts:
        return pd.DataFrame(columns=["year", "company_name_norm", "sector_description", "industry_description"])
    combo = pd.concat(parts, ignore_index=True, sort=False)
    combo = (
        combo.groupby(["year", "company_name_norm"], dropna=False)
        .agg(
            sector_description=("sector_description", _first_non_empty),
            industry_description=("industry_description", _first_non_empty),
        )
        .reset_index()
    )
    combo["year"] = combo["year"].astype(int)
    return combo


def load_lseg_presence(lseg_root: Path) -> pd.DataFrame:
    ten_year = lseg_root / "lseg_active_universe_10y.csv"
    parts: List[pd.DataFrame] = []

    if ten_year.exists():
        df = pd.read_csv(ten_year, low_memory=False, usecols=lambda c: c in {"Instrument", "Company Common Name", "ISIN", "FiscalYear"})
        parts.append(df)
    else:
        for path in sorted(lseg_root.glob("lseg_active_universe_*.csv")):
            df = pd.read_csv(path, low_memory=False, usecols=lambda c: c in {"Instrument", "Company Common Name", "ISIN", "FiscalYear"})
            parts.append(df)

    if not parts:
        return pd.DataFrame(columns=["year", "lseg_instrument", "lseg_company_name", "lseg_isin", "lseg_name_norm"])

    lseg = pd.concat(parts, ignore_index=True, sort=False)
    lseg["year"] = lseg.get("FiscalYear", pd.Series(dtype=object)).map(_clean_year)
    lseg["lseg_instrument"] = lseg.get("Instrument", pd.Series(dtype=object)).fillna("").astype(str).str.strip()
    lseg["lseg_company_name"] = lseg.get("Company Common Name", pd.Series(dtype=object)).fillna("").astype(str).str.strip()
    lseg["lseg_isin"] = lseg.get("ISIN", pd.Series(dtype=object)).map(_clean_isin)
    lseg["lseg_name_norm"] = lseg["lseg_company_name"].map(normalize_name)
    lseg = lseg[(lseg["year"].notna()) & ((lseg["lseg_isin"] != "") | (lseg["lseg_name_norm"] != ""))]
    return lseg[["year", "lseg_instrument", "lseg_company_name", "lseg_isin", "lseg_name_norm"]].drop_duplicates()


def load_trucost_emissions(trucost_root: Path) -> pd.DataFrame:
    ghg_path = trucost_root / "260710 trucost pulic-ghg-2011to24.dta"
    if not ghg_path.exists():
        raise FileNotFoundError(f"Trucost emissions file not found: {ghg_path}")

    hdr = pd.read_stata(ghg_path, nrows=0)
    di_cols = [c for c in hdr.columns if c.lower().startswith("di_")]
    keep_cols = [c for c in ["companyid", "fiscalyear", "companyname", "ticker", "gvkey"] if c in hdr.columns] + di_cols
    df = pd.read_stata(ghg_path, usecols=keep_cols)

    df["year"] = df["fiscalyear"].map(_clean_year)
    df["trucost_companyid_key"] = df["companyid"].map(_clean_company_id)
    df["trucost_companyname_emissions"] = df["companyname"].fillna("").astype(str).str.strip()
    df["trucost_ticker_emissions"] = df["ticker"].fillna("").astype(str).str.strip()
    df = df[(df["year"].notna()) & (df["trucost_companyid_key"] != "")]

    out_cols = ["year", "trucost_companyid_key", "trucost_companyname_emissions", "trucost_ticker_emissions"] + di_cols
    return df[out_cols].drop_duplicates(subset=["year", "trucost_companyid_key"])


def build_common_four_with_trucost_emissions(merged_df: pd.DataFrame, lseg_root: Path, trucost_root: Path) -> pd.DataFrame:
    base = merged_df.copy()
    base["year"] = base["trucost_fiscalyear"].map(_clean_year)
    base["trucost_isin_key"] = base["trucost_isin"].map(_clean_isin)
    base["trucost_name_norm"] = base["trucost_companyname"].map(normalize_name)
    base["trucost_companyid_key"] = base["trucost_companyid"].map(_clean_company_id)
    base["cdp_name_present"] = base["cdp_org_name"].fillna("").astype(str).str.strip() != ""
    base["factset_name_present"] = base["factset_company_name"].fillna("").astype(str).str.strip() != ""
    base = base[(base["year"].notna()) & base["cdp_name_present"] & base["factset_name_present"]].copy()

    lseg = load_lseg_presence(lseg_root)
    if lseg.empty:
        raise RuntimeError(f"No LSEG yearly files found in {lseg_root}")

    lseg_isin = lseg[lseg["lseg_isin"] != ""].copy()
    isin_merge = base.merge(
        lseg_isin,
        left_on=["year", "trucost_isin_key"],
        right_on=["year", "lseg_isin"],
        how="left",
    )

    missing = isin_merge["lseg_instrument"].fillna("").astype(str).str.strip() == ""
    if missing.any():
        lseg_name = lseg[lseg["lseg_name_norm"] != ""].copy()
        by_name = isin_merge.loc[missing, ["year", "trucost_name_norm"]].merge(
            lseg_name[["year", "lseg_name_norm", "lseg_instrument", "lseg_company_name", "lseg_isin"]],
            left_on=["year", "trucost_name_norm"],
            right_on=["year", "lseg_name_norm"],
            how="left",
        )
        isin_merge.loc[missing, "lseg_instrument"] = by_name["lseg_instrument"].values
        isin_merge.loc[missing, "lseg_company_name"] = by_name["lseg_company_name"].values
        isin_merge.loc[missing, "lseg_isin"] = by_name["lseg_isin"].values

    common = isin_merge[isin_merge["lseg_instrument"].fillna("").astype(str).str.strip() != ""].copy()
    emissions = load_trucost_emissions(trucost_root)
    common = common.merge(emissions, on=["year", "trucost_companyid_key"], how="left")

    preferred_cols = [
        "year",
        "trucost_companyid",
        "trucost_ticker",
        "trucost_companyname",
        "trucost_isin",
        "factset_entity_id",
        "factset_company_name",
        "factset_ticker",
        "cdp_org_name",
        "cdp_match_score",
        "match_type",
        "fuzzy_score",
        "lseg_instrument",
        "lseg_company_name",
        "lseg_isin",
    ]
    di_cols = [c for c in common.columns if c.lower().startswith("di_")]
    final_cols = [c for c in preferred_cols if c in common.columns] + di_cols
    out = common[final_cols].copy()
    out = out.drop_duplicates(subset=[c for c in ["year", "trucost_companyid", "trucost_companyname"] if c in out.columns])
    out = out.sort_values([c for c in ["year", "trucost_companyname"] if c in out.columns])
    return out


def build_unique_company_universe_by_year(merged_df: pd.DataFrame, cdp_root: Path) -> pd.DataFrame:
    year_series = merged_df.get("trucost_fiscalyear", pd.Series(dtype=float)).map(_clean_year)
    years = sorted(set(y for y in year_series.dropna().tolist() if y is not None))
    if not years:
        years = list(range(2014, 2026))

    t = merged_df.copy()
    if "simpleindustry" not in t.columns:
        t["simpleindustry"] = ""
    if "tcprimarysectorid" not in t.columns:
        t["tcprimarysectorid"] = ""
    t = t[["trucost_fiscalyear", "trucost_companyname", "simpleindustry", "tcprimarysectorid"]].copy()
    t["year"] = t["trucost_fiscalyear"].map(_clean_year)
    t["source"] = "trucost"
    t["company_name"] = t["trucost_companyname"].fillna("").astype(str).str.strip()
    t["sector_description"] = t["tcprimarysectorid"].fillna("").astype(str).str.strip()
    t["industry_description"] = t["simpleindustry"].fillna("").astype(str).str.strip()
    t = t[(t["year"].notna()) & (t["company_name"] != "")]
    t["company_name_norm"] = t["company_name"].map(normalize_name)
    trucost_desc = load_trucost_descriptions(TRUCOST_RAW_ROOT)
    t = t.merge(trucost_desc, on=["year", "company_name_norm"], how="left", suffixes=("", "_raw"))
    t["sector_description"] = t["sector_description"].where(t["sector_description"] != "", t["sector_description_raw"].fillna(""))
    t["industry_description"] = t["industry_description"].where(t["industry_description"] != "", t["industry_description_raw"].fillna(""))
    trucost_unique = t[["year", "source", "company_name", "company_name_norm", "sector_description", "industry_description"]].drop_duplicates()

    fs_desc = load_factset_entity_descriptions(FACTSET_RAW_ROOT)
    f = merged_df[["trucost_fiscalyear", "factset_entity_id", "factset_company_name"]].copy()
    f["year"] = f["trucost_fiscalyear"].map(_clean_year)
    f["source"] = "factset"
    f["factset_entity_id"] = f["factset_entity_id"].fillna("").astype(str).str.strip()
    f["company_name"] = f["factset_company_name"].fillna("").astype(str).str.strip()
    f = f[(f["year"].notna()) & (f["company_name"] != "")]
    f = f.merge(fs_desc, left_on="factset_entity_id", right_on="FACTSET_ENTITY_ID", how="left")
    f["sector_description"] = f["sector_description"].fillna("").astype(str).str.strip()
    f["industry_description"] = f["industry_description"].fillna("").astype(str).str.strip()
    f["company_name_norm"] = f["company_name"].map(normalize_name)
    factset_unique = f[["year", "source", "company_name", "company_name_norm", "sector_description", "industry_description"]].drop_duplicates()

    cdp_unique = load_cdp_unique_names_by_year(cdp_root, years)

    combined = pd.concat([trucost_unique, factset_unique, cdp_unique], ignore_index=True, sort=False)
    combined["year"] = combined["year"].astype(int)
    combined["sector_description"] = combined["sector_description"].fillna("").astype(str).str.strip()
    combined["industry_description"] = combined["industry_description"].fillna("").astype(str).str.strip()
    combined = (
        combined.groupby(["year", "source", "company_name_norm"], dropna=False)
        .agg(
            company_name=("company_name", _first_non_empty),
            sector_description=("sector_description", _first_non_empty),
            industry_description=("industry_description", _first_non_empty),
        )
        .reset_index()
        .sort_values(["year", "source", "company_name"])
    )
    return combined


def main() -> None:
    parser = argparse.ArgumentParser(description="Combined merge pipeline: Trucost/FactSet rematch + CDP score refresh + yearly unique name universe.")
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output merged CSV path (default: overwrite input)",
    )
    parser.add_argument(
        "--cdp-root",
        type=Path,
        default=Path(str(DEFAULT_CDP_ROOT)),
        help="CDP root folder for yearly unique-name extraction",
    )
    parser.add_argument(
        "--unique-output",
        type=Path,
        default=Path(r"/data/processed/unique_company_names_by_year_trucost_factset_lseg_cdp.csv"),
        help="Output CSV for unique company names by source and year",
    )
    parser.add_argument(
        "--lseg-root",
        type=Path,
        default=LSEG_RAW_ROOT,
        help="Folder containing yearly LSEG active universe CSVs",
    )
    parser.add_argument(
        "--common-four-output",
        type=Path,
        default=Path(r"/data/processed/common_companies_by_year_cdp_trucost_factset_lseg_with_trucost_emissions.csv"),
        help="Output CSV for companies present in CDP, Trucost, FactSet and LSEG by year, with Trucost emissions columns",
    )
    args = parser.parse_args()

    input_path = args.input
    if not input_path.exists():
        alt = Path(r"C:\Users\scelik\Desktop\TRACE3Code\data\outputs\dataset_readers\common_factset_trucost_cdp_tickers_extended.csv")
        if alt.exists():
            input_path = alt
        else:
            raise FileNotFoundError(f"Input file not found: {args.input}")
    output_path = args.output if args.output is not None else input_path

    df = pd.read_csv(input_path, low_memory=False)

    old_factset = df["factset_company_name"].astype(str)
    old_cdp_score = pd.to_numeric(df["cdp_match_score"], errors="coerce")

    rematched = rematch_trucost_factset(df)
    updated = update_cdp_match_scores(rematched)
    unique_by_year = build_unique_company_universe_by_year(updated, args.cdp_root)
    common_four = build_common_four_with_trucost_emissions(updated, args.lseg_root, TRUCOST_RAW_ROOT)

    new_factset = updated["factset_company_name"].astype(str)
    new_cdp_score = pd.to_numeric(updated["cdp_match_score"], errors="coerce")
    changed_factset = int((old_factset != new_factset).sum())
    changed_cdp = int((old_cdp_score.fillna(-1) != new_cdp_score.fillna(-1)).sum())
    non_eq_full = int((((updated["trucost_companyname"].map(normalize_name) != updated["factset_company_name"].map(normalize_name)) & pd.to_numeric(updated["fuzzy_score"], errors="coerce").fillna(0).eq(1000)).sum()))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    updated.to_csv(output_path, index=False)

    args.unique_output.parent.mkdir(parents=True, exist_ok=True)
    unique_by_year.to_csv(args.unique_output, index=False)
    args.common_four_output.parent.mkdir(parents=True, exist_ok=True)
    common_four.to_csv(args.common_four_output, index=False)

    print(f"Saved merged: {output_path}")
    print(f"Merged rows: {len(updated)}")
    print(f"Changed FactSet match rows: {changed_factset}")
    print(f"Changed cdp_match_score rows: {changed_cdp}")
    print(f"Non-equivalent Trucost/FactSet rows with score=1000: {non_eq_full}")
    print(f"Saved unique names by year: {args.unique_output}")
    print(f"Unique rows: {len(unique_by_year)}")
    print(f"Saved common-4 yearly file with Trucost emissions: {args.common_four_output}")
    print(f"Common-4 rows: {len(common_four)}")


if __name__ == "__main__":
    main()
