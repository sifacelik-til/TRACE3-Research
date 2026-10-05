"""Extract and cluster CDP 2016–2024 climate initiatives, risks, opportunities.

Run::

    python -m src.cdp_text_clustering.cluster_cdp_2016_2024_climate_actions --stage extract
    python -m src.cdp_text_clustering.cluster_cdp_2016_2024_climate_actions --stage enrich
    python -m src.cdp_text_clustering.cluster_cdp_2016_2024_climate_actions --stage baseline
    python -m src.cdp_text_clustering.cluster_cdp_2016_2024_climate_actions --stage cluster

The first stage needs openpyxl, pandas, pypdf and pyarrow; it does not import
PyTorch. The optional BERTopic stage uses a local multilingual model. No CDP
responses are sent to an external service. Raw reported fields, quantitative
estimates, and conservative text-derived signals have separate columns.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import html
import logging
import os
import re
import sys
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

# On Windows, c10.dll can fail if numerical libraries initialize first. Keep
# extraction torch-free, but load torch before NumPy/pandas for clustering.
if any(arg in {"cluster", "all"} for arg in sys.argv[1:]):
    numba_cache = Path(tempfile.gettempdir()) / "codex_cdp_numba_cache"
    numba_cache.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("NUMBA_CACHE_DIR", str(numba_cache))
    os.environ.setdefault("NUMBA_NUM_THREADS", "4")
    os.environ.setdefault("HF_HUB_OFFLINE", "1")
    import torch
    torch.set_num_threads(min(8, torch.get_num_threads()))

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
INPUTS = {
    2016: ROOT / "data/raw/CDP/2016/CDP_2016_ClimateChange_Public.xlsx",
    2017: ROOT / "data/raw/CDP/2017/CDP_2017_ClimateChange_Public.xlsx",
    2018: ROOT / "data/raw/CDP/2018/CDP_2018_ClimateChange_Inv_SC_Public_Tabbed.xlsx",
    2019: ROOT / "data/raw/CDP/2019/CDP_2019_ClimateChange_Inv_SC_Public_Tabbed.xlsx",
    2020: ROOT / "data/raw/CDP/2020/CDP_2020_ClimateChange_Public_Inv_SC.xlsx",
    2021: ROOT / "data/raw/CDP/2021/CDP_2021_ClimateChange_Public_Inv_SC.xlsx",
    2022: ROOT / "data/raw/CDP/2022/CDP_2022_ClimateChange_Public_Inv_SC_v2.2.xlsx",
    2023: ROOT / "data/raw/CDP/2023/CDP_2023_ClimateChange_Public_Inv_SCv3.xlsx",
}
GUIDES = {
    year: ROOT / f"data/raw/CDP/CDP Questionnaires/{year} Climate Change Questionnaire.pdf"
    for year in INPUTS
}
GUIDES[2024] = {
    "Q7.55.2": ROOT / "data/raw/CDP/CDP Questionnaires/2024 - Corporate Questionnaire - Modules 7.pdf",
    "Q3.1.1": ROOT / "data/raw/CDP/CDP Questionnaires/2024 - Corporate Questionnaire - Modules 1 to 6.pdf",
    "Q3.6.1": ROOT / "data/raw/CDP/CDP Questionnaires/2024 - Corporate Questionnaire - Modules 1 to 6.pdf",
}
# The "noisin" response file carries the same question rows and organization
# universe; it is an alternate identifier packaging, not an additional sample.
INPUT_2024 = ROOT / "data/raw/CDP/2024/full_extract_cm_eds_c_isin_2024_responses_v1_20250623_125717.parquet"
OUTPUT = ROOT / "data/processed/cdp_2016_2024_climate_actions"
RECORDS = OUTPUT / "climate_action_risk_opportunity_records.csv.gz"
MODEL_NAME = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
MODERN_SHEETS = {"C4.3b": "initiative", "C2.3a": "risk", "C2.4a": "opportunity"}
LEGACY_SHEETS = {
    "CC3.3b": "initiative",
    "CC5.1a": "risk", "CC5.1b": "risk", "CC5.1c": "risk",
    "CC6.1a": "opportunity", "CC6.1b": "opportunity", "CC6.1c": "opportunity",
}
PARQUET_QUESTIONS = {"Q7.55.2": "initiative", "Q3.1.1": "risk", "Q3.6.1": "opportunity"}
GUIDE_PAGES = {
    2016: {"CC3.3b": 8, "CC5.1a": 9, "CC5.1b": 9, "CC5.1c": 9, "CC6.1a": 9, "CC6.1b": 9, "CC6.1c": 9},
    2017: {"CC3.3b": 7, "CC5.1a": 8, "CC5.1b": 8, "CC5.1c": 8, "CC6.1a": 8, "CC6.1b": 8, "CC6.1c": 8},
    2018: {"C4.3b": 50, "C2.3a": 32, "C2.4a": 34},
    2019: {"C4.3b": 57, "C2.3a": 35, "C2.4a": 39},
    2020: {"C4.3b": 57, "C2.3a": 31, "C2.4a": 34},
    2021: {"C4.3b": 39, "C2.3a": 16, "C2.4a": 18},
    2022: {"C4.3b": 70, "C2.3a": 38, "C2.4a": 41},
    2023: {"C4.3b": 73, "C2.3a": 36, "C2.4a": 39},
    2024: {"Q7.55.2": 322, "Q3.1.1": 127, "Q3.6.1": 167},
}
MISSING = {"", "question not applicable", "not applicable", "n/a", "nan", "none"}


@dataclass
class CurrencyRegistry:
    values: dict[tuple[int, str], tuple[str, str, str, str]]
    unavailable_years: dict[int, str]
    ambiguous: set[tuple[int, str]]


def clean(value: object) -> str:
    if value is None:
        return ""
    result = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.unescape(str(value)))).strip()
    return "" if result.casefold() in MISSING else result


def number(value: object) -> float:
    text = clean(value).replace(",", "")
    try:
        result = float(text)
        return result if np.isfinite(result) else np.nan
    except ValueError:
        return np.nan


def normalize_label(value: object) -> str:
    """Standardize typography lost by some console/PDF encodings."""
    return re.sub(r"[^a-z0-9]+", " ", clean(value).casefold()).strip()


def account_key(value: object) -> str:
    """Use the same account identifier for workbook and CSV/Parquet records."""
    account = clean(value)
    return account[:-2] if re.fullmatch(r"\d+\.0", account) else account


def currency_iso(raw: str) -> str:
    """Conservatively parse CDP's three-letter currency selection."""
    match = re.match(r"^\s*([A-Za-z]{3})(?=\s|\(|$)", raw)
    return match.group(1).upper() if match else ""


def load_reporting_currencies() -> CurrencyRegistry:
    """Read company-year currency selections from CC0.4/C0.4 and 2024 Q1.2.

    A temporarily inaccessible year remains explicitly unpopulated. Conflicting
    values for a company-year are marked ambiguous instead of choosing one.
    """
    registry = CurrencyRegistry({}, {}, set())

    def register(year: int, account: object, raw_value: object, source_file: Path,
                 source_sheet: str, question: str) -> None:
        key = (year, account_key(account))
        raw = clean(raw_value)
        if not key[1] or not raw or key in registry.ambiguous:
            return
        value = (raw, str(source_file.relative_to(ROOT)), source_sheet, question)
        previous = registry.values.get(key)
        if previous and previous[0] != raw:
            registry.values.pop(key, None)
            registry.ambiguous.add(key)
        else:
            registry.values[key] = value

    for year, path in INPUTS.items():
        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except (OSError, PermissionError) as exc:
            registry.unavailable_years[year] = f"{type(exc).__name__}: {exc}"
            continue
        try:
            sheet = next(
                (name for name in ("C0.4", "C0 - Introduction", "CC0. Introduction")
                 if name in book.sheetnames),
                None,
            )
            if sheet is None:
                registry.unavailable_years[year] = "Currency question sheet is absent"
                continue
            rows = book[sheet].iter_rows(values_only=True)
            for header in rows:
                labels = [clean(value).casefold() for value in header]
                account_col = next(
                    (i for i, label in enumerate(labels)
                     if label in {"account number", "account_id"}),
                    None,
                )
                currency_col = next(
                    (i for i, label in enumerate(labels)
                     if ("cc0.4" in label or "c0.4" in label)
                     and "currency" in label),
                    None,
                )
                if account_col is not None and currency_col is not None:
                    break
            else:
                registry.unavailable_years[year] = f"Currency header not found in {sheet}"
                continue
            question = "CC0.4" if year <= 2017 else "C0.4"
            for row in rows:
                register(year, row[account_col], row[currency_col], path, sheet, question)
        finally:
            book.close()

    try:
        import pyarrow.dataset as ds

        dataset = ds.dataset(INPUT_2024, format="parquet")
        selected = dataset.scanner(
            columns=["cdp_disclosing_org_number", "column_header",
                     "content_full", "question_number"],
            filter=(ds.field("question_number") == "Q1.2")
                   & (ds.field("public_status") == "Public"),
            batch_size=40_000,
        )
        for batch in selected.to_batches():
            for cell in batch.to_pylist():
                if "select the currency used" not in clean(cell["column_header"]).casefold():
                    continue
                register(
                    2024, cell["cdp_disclosing_org_number"], cell["content_full"],
                    INPUT_2024, "Q1.2", "Q1.2",
                )
    except (OSError, PermissionError) as exc:
        registry.unavailable_years[2024] = f"{type(exc).__name__}: {exc}"
    return registry


def currency_fields(registry: CurrencyRegistry, year: int, account: object) -> dict:
    key = (year, account_key(account))
    value = registry.values.get(key)
    if value:
        raw, source_file, source_sheet, question = value
        code = currency_iso(raw)
        status = "matched" if code else "unrecognized_currency_code"
    else:
        raw = code = source_file = source_sheet = question = ""
        status = (
            "ambiguous" if key in registry.ambiguous else
            "source_unavailable" if year in registry.unavailable_years else
            "not_reported"
        )
    return {
        "reporting_currency_raw": raw,
        "reporting_currency_iso": code,
        "reporting_currency_source_file": source_file,
        "reporting_currency_source_sheet": source_sheet,
        "reporting_currency_source_question": question,
        "reporting_currency_status": status,
        "monetary_unit": code,
    }


def field_headers(headers: tuple[object, ...]) -> dict[str, int]:
    """Map the CDP label after ' - '; column positions differ by year."""
    labels = {}
    for index, value in enumerate(headers):
        if not value:
            continue
        header = clean(value)
        # Legacy CC headers contain two question/field separators. The field
        # itself can also contain " - ", e.g. "unit currency - as specified".
        split_count = 2 if re.match(r"^CC\S+ C\d+ - ", header) else 1
        labels[normalize_label(header.split(" - ", split_count)[-1])] = index
    return labels


def get(row: tuple[object, ...], columns: dict[str, int], *aliases: str) -> str:
    for alias in aliases:
        index = columns.get(normalize_label(alias))
        if index is not None:
            return clean(row[index])
    return ""


def get_numeric(row: tuple[object, ...], columns: dict[str, int], *aliases: str) -> float:
    return number(get(row, columns, *aliases))


def get_prefix(row: tuple[object, ...], columns: dict[str, int], prefix: str) -> str:
    key = normalize_label(prefix)
    matches = [index for label, index in columns.items() if label.startswith(key)]
    if len(matches) != 1:
        return ""
    return clean(row[matches[0]])


CONDITION_PATTERN = re.compile(
    r"\b(?:depend(?:s|ent|ing)? on|requires?|provided that|subject to|"
    r"conditional on|only if|necessary to|in order to|prerequisite|contingent on)\b",
    re.I,
)
COMPLEXITY_PATTERNS = {
    "high": re.compile(r"\b(?:highly complex|significant technical challenge|major implementation challenge|difficult to implement|substantial operational complexity)\b", re.I),
    "medium": re.compile(r"\b(?:moderately complex|some implementation challenge|requires coordination across|requires significant coordination)\b", re.I),
    "low": re.compile(r"\b(?:straightforward to implement|simple to implement|low operational complexity|minimal operational disruption)\b", re.I),
}


def evidence_sentence(text: str, pattern: re.Pattern[str]) -> str:
    for sentence in re.split(r"(?<=[.!?])\s+|\n+", text):
        if pattern.search(sentence):
            return sentence[:600].strip()
    return ""


def derive_signals(text: str) -> tuple[str, str, str]:
    condition = evidence_sentence(text, CONDITION_PATTERN)
    for label, pattern in COMPLEXITY_PATTERNS.items():
        phrase = evidence_sentence(text, pattern)
        if phrase:
            return label, phrase, condition
    return "", "", condition


def common(row: tuple[object, ...], columns: dict[str, int], year: int, sheet: str,
           excel_row: int, currencies: CurrencyRegistry) -> dict:
    record_type = (LEGACY_SHEETS if year <= 2017 else MODERN_SHEETS)[sheet]
    account = get(row, columns, "Account number", "account_id")
    return {
        "year": year,
        "year_definition": "CDP questionnaire/disclosure year; fiscal reporting year may differ",
        "reporting_year_reported": get(row, columns, "accounting_year"),
        "cdp_account_number": account,
        "organization": get(row, columns, "Organization", "account_name"),
        "country": get(row, columns, "Country", "Country/Area"),
        "primary_sector": get(row, columns, "Primary sector"),
        "primary_industry": get(row, columns, "Primary industry"),
        "question_id": sheet,
        "record_type": record_type,
        "question_row": get(row, columns, "Row"),
        "source_excel_row": excel_row,
        "source_file": str(INPUTS[year].relative_to(ROOT)),
        "source_sheet": sheet,
        "guide_file": str(GUIDES[year].relative_to(ROOT)),
        **currency_fields(currencies, year, account),
    }


def initiative(row: tuple[object, ...], c: dict[str, int], result: dict) -> None:
    year = result["year"]
    category = get(row, c, "Initiative category & Initiative type_G", "Initiative type", "Activity type")
    initiative_type = get(row, c, "Initiative category & Initiative type", "Description of initiative", "Description of activity")
    if year == 2024 and ":" in initiative_type:
        category, initiative_type = [part.strip() for part in initiative_type.split(":", 1)]
    comment = get(row, c, "Comment")
    result.update({
        "initiative_category_reported": category,
        "initiative_type_reported": initiative_type,
        "initiative_comment_reported": comment,
        "reported_outcome": "Estimated annual CO2e savings (not measured realized reduction)",
        "reported_outcome_text": comment,
        "emissions_boundary_reported": get(row, c, "Scope", "Scope(s)", "Scope(s) or Scope 3 category(ies) where emissions savings occur"),
        "estimated_annual_emissions_impact_tco2e": get_numeric(row, c, "Estimated annual CO2e savings (metric tonnes CO2e)", "Estimated annual CO2e savings (metric tons CO2e)"),
        "annual_monetary_savings_reported": number(get_prefix(row, c, "Annual monetary savings")),
        "cost_reported": number(get_prefix(row, c, "Investment required")),
        "cost_text_reported": get_prefix(row, c, "Investment required"),
        # "cost_kind": "Investment required",
        "payback_period_reported": get(row, c, "Payback period"),
        "initiative_lifetime_reported": get(row, c, "Estimated lifetime of the initiative"),
        "implementation_time_reported": "Implemented in reporting year; precise duration unavailable",
        "voluntary_or_mandatory": get(row, c, "Voluntary/Mandatory", "Voluntary/ Mandatory"),
        "source_or_driver_reported": initiative_type or category,
        "financial_impact_reported": np.nan,
        "financial_impact_kind": "",
        "financial_impact_description": "",
        "timeframe_reported": "",
        "timeframe_kind": "",
        "value_chain_position_reported": "",
        "response_strategy_reported": "",
        "narrative_for_clustering": " ".join(filter(None, [category, initiative_type, comment])),
    })


def risk_or_opportunity(row: tuple[object, ...], c: dict[str, int], result: dict) -> None:
    is_risk = result["record_type"] == "risk"
    description = get(row, c, "Company-specific description", "Company- specific description", "Description")
    strategy = get(row, c, "Management method", "Description of response and explanation of cost calculation") if is_risk else get(row, c, "Strategy to realize opportunity", "Strategy to realize opportunity and explanation of cost calculation", "Management method")
    driver = get(row, c, "Primary climate-related risk driver", "Risk type & Primary climate-related risk driver", "Primary climate-related opportunity driver", "Risk driver", "Opportunity driver")
    driver_group = get(row, c, "Risk type & Primary climate-related risk driver_G", "Risk type", "Opportunity type")
    if result["year"] <= 2017:
        driver_group = {"a": "regulatory", "b": "physical", "c": "other climate-related"}[result["question_id"][-1]]
    cost_text = get(row, c, "Cost of management", "Cost of response to risk", "Cost to realize opportunity")
    financial_text = get(row, c, "Estimated financial implications", "Potential financial impact")
    result.update({
        "initiative_category_reported": "",
        "initiative_type_reported": "",
        "initiative_comment_reported": "",
        "reported_outcome": "Potential financial or strategic impact, not observed emissions reduction",
        "reported_outcome_text": description,
        "emissions_boundary_reported": "",  # Value chain position is not a GHG Protocol scope.
        "estimated_annual_emissions_impact_tco2e": np.nan,
        "annual_monetary_savings_reported": np.nan,
        "cost_reported": number(cost_text),
        "cost_text_reported": cost_text,
        "cost_kind": "Cost of response to risk" if is_risk else "Cost to realize opportunity",
        "payback_period_reported": "",
        "initiative_lifetime_reported": "",
        "implementation_time_reported": "",  # Risk horizon is not implementation time.
        "voluntary_or_mandatory": "",
        "source_or_driver_reported": driver,
        "source_or_driver_group": driver_group,
        "financial_impact_reported": get_numeric(row, c, "Potential financial impact figure (currency)") if not financial_text else number(financial_text),
        "financial_impact_min_reported": get_numeric(row, c, "Potential financial impact figure � minimum (currency)"),
        "financial_impact_max_reported": get_numeric(row, c, "Potential financial impact figure � maximum (currency)"),
        "financial_impact_kind": get(row, c, "Type of financial impact", "Type of financial impact driver", "Primary potential financial impact", "Potential impact"),
        "financial_impact_description": " | ".join(filter(None, [financial_text if np.isnan(number(financial_text)) else "", get(row, c, "Explanation of financial impact figure", "Explanation of financial impact")])),
        "timeframe_reported": get(row, c, "Time horizon", "Timeframe"),
        "timeframe_kind": "Risk/opportunity time horizon",
        "value_chain_position_reported": get(row, c, "Where in the value chain does the risk driver occur?", "Where in the value chain does the opportunity occur?", "Direct/Indirect", "Direct/ Indirect"),
        "response_strategy_reported": strategy,
        "likelihood_reported": get(row, c, "Likelihood"),
        "magnitude_reported": get(row, c, "Magnitude of impact"),
        "narrative_for_clustering": " ".join(filter(None, [driver_group, driver, description, strategy])),
    })


def glossary() -> pd.DataFrame:
    """Verify every mapped question in its own-year questionnaire PDF."""
    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)
    rows = []
    cache = {}
    for year, questions in GUIDE_PAGES.items():
        for question, page in questions.items():
            path = GUIDES[year][question] if year == 2024 else GUIDES[year]
            if not path.exists():
                raise FileNotFoundError(path)
            if path not in cache:
                cache[path] = PdfReader(path)
            text = cache[path].pages[page - 1].extract_text() or ""
            guide_id = (
                question[:-1] if year <= 2017 and question.startswith(("CC5.1", "CC6.1"))
                else question.removeprefix("Q")
            )
            if guide_id not in text:
                raise ValueError(f"Could not verify {year} {question} (via {guide_id}) in {path} p.{page}")
            prompt_match = re.search(rf"\(?{re.escape(guide_id)}\)?\s*([^\n]{{10,250}})", text)
            rows.append({
                "year": year,
                "question_id": question,
                "questionnaire_section": guide_id,
                "question_prompt_excerpt": clean(prompt_match.group(1)) if prompt_match else "",
                "guide_file": str(path.relative_to(ROOT)),
                "pdf_page": page,
                "verification": "Question ID found on the mapped questionnaire page",
            })
    return pd.DataFrame(rows)


OUTPUT_FIELDS = [
    "record_id", "year", "year_definition", "reporting_year_reported",
    "cdp_account_number", "organization", "country", "primary_sector",
    "primary_industry", "question_id", "record_type", "question_row",
    "source_excel_row", "source_file", "source_sheet", "guide_file",
    "monetary_unit", "reporting_currency_raw", "reporting_currency_iso",
    "reporting_currency_source_file", "reporting_currency_source_sheet",
    "reporting_currency_source_question", "reporting_currency_status",
    "initiative_category_reported", "initiative_type_reported",
    "initiative_comment_reported",
    "reported_outcome", "reported_outcome_text", "emissions_boundary_reported",
    "estimated_annual_emissions_impact_tco2e", "annual_monetary_savings_reported",
    "cost_reported", "cost_text_reported", "cost_kind", "payback_period_reported",
    "initiative_lifetime_reported", "implementation_time_reported",
    "voluntary_or_mandatory", "source_or_driver_reported", "source_or_driver_group",
    "financial_impact_reported", "financial_impact_min_reported",
    "financial_impact_max_reported", "financial_impact_kind",
    "financial_impact_description", "financial_effect_reporting_year_reported",
    "anticipated_financial_effect_short_min_reported",
    "anticipated_financial_effect_short_max_reported",
    "anticipated_financial_effect_medium_min_reported",
    "anticipated_financial_effect_medium_max_reported",
    "anticipated_financial_effect_long_min_reported",
    "anticipated_financial_effect_long_max_reported",
    "timeframe_reported", "timeframe_kind", "value_chain_position_reported",
    "response_strategy_reported", "likelihood_reported", "magnitude_reported",
    "narrative_for_clustering", "operational_complexity_signal_inferred",
    "operational_complexity_evidence", "condition_for_success_evidence_inferred",
    "text_signal_method",
]


def complete_record(entry: dict) -> dict:
    """Keep heuristic signals separate from questionnaire responses."""
    signal_text = (
        entry.get("reported_outcome_text", "")
        if entry["record_type"] == "initiative"
        else entry.get("response_strategy_reported", "")
    )
    complexity, complexity_evidence, condition_evidence = derive_signals(signal_text)
    entry["operational_complexity_signal_inferred"] = complexity
    entry["operational_complexity_evidence"] = complexity_evidence
    entry["condition_for_success_evidence_inferred"] = condition_evidence
    entry["text_signal_method"] = "Conservative keyword + sentence evidence; blank means not identified"
    return entry


def write_entry(writer: csv.DictWriter, entry: dict) -> None:
    unexpected = set(entry) - set(OUTPUT_FIELDS)
    if unexpected:
        raise ValueError(f"Unmapped output columns: {sorted(unexpected)}")
    writer.writerow({key: "" if pd.isna(value) else value for key, value in entry.items()})


def risk_or_opportunity_2024(row: tuple[object, ...], c: dict[str, int], result: dict) -> None:
    """Keep 2024 current and anticipated financial effects in distinct fields."""
    is_risk = result["record_type"] == "risk"
    kind = "risk" if is_risk else "opportunity"
    driver = get(row, c, f"{kind.title()} types and primary environmental {kind} driver", f"{kind.title()} type and primary environmental {kind} driver")
    description = get(row, c, "Organization-specific description of risk", "Organization specific description")
    strategy = get(row, c, "Description of response", "Strategy to realize opportunity")
    cost_text = get(row, c, "Cost of response to risk", "Cost to realize opportunity")
    current_effect = get_numeric(row, c, "Financial effect figure in the reporting year (currency)")
    result.update({
        "initiative_category_reported": "",
        "initiative_type_reported": "",
        "initiative_comment_reported": "",
        "reported_outcome": "Current or anticipated financial/strategic effect, not observed emissions reduction",
        "reported_outcome_text": description,
        "emissions_boundary_reported": "",
        "estimated_annual_emissions_impact_tco2e": np.nan,
        "annual_monetary_savings_reported": np.nan,
        "cost_reported": number(cost_text),
        "cost_text_reported": cost_text,
        "cost_kind": "Cost of response to risk" if is_risk else "Cost to realize opportunity",
        "payback_period_reported": "",
        "initiative_lifetime_reported": "",
        "implementation_time_reported": "",
        "voluntary_or_mandatory": "",
        "source_or_driver_reported": driver,
        "source_or_driver_group": driver.split(":", 1)[0].strip() if ":" in driver else "",
        "financial_impact_reported": current_effect,
        "financial_effect_reporting_year_reported": current_effect,
        "financial_impact_kind": get(row, c, f"Primary financial effect of the {kind}"),
        "financial_impact_description": get(row, c, "Explanation of financial effect figure", "Explanation of financial effect figures"),
        "timeframe_reported": get(row, c, f"Time horizon over which the {kind} is anticipated to have a substantive effect on the organization"),
        "timeframe_kind": "Risk/opportunity effect horizon",
        "value_chain_position_reported": get(row, c, f"Value chain stage where the {kind} occurs"),
        "response_strategy_reported": strategy,
        "likelihood_reported": get(row, c, f"Likelihood of the {kind} having an effect within the anticipated time horizon"),
        "magnitude_reported": get(row, c, "Magnitude"),
        "narrative_for_clustering": " ".join(filter(None, [driver, description, strategy])),
    })
    # The 2024 questionnaire distinguishes three anticipated horizons. Do not
    # collapse these to a single potential impact figure.
    for horizon in ("short", "medium", "long"):
        for bound in ("minimum", "maximum"):
            value = get_numeric(row, c, f"Anticipated financial effect figure in the {horizon}-term - {bound} (currency)")
            result[f"anticipated_financial_effect_{horizon}_{'min' if bound == 'minimum' else 'max'}_reported"] = value


def extract_2024(writer: csv.DictWriter, audit: list[dict],
                 max_rows_per_question: int | None, currencies: CurrencyRegistry) -> int:
    """Stream public climate response cells from the 2024 Parquet extract."""
    import pyarrow.dataset as ds

    total = 0
    for path in [INPUT_2024]:
        if not path.exists():
            raise FileNotFoundError(path)
        dataset = ds.dataset(path, format="parquet")
        source_name = "isin"
        columns = [
            "cdp_disclosing_org_number", "disclosing_organization",
            "discloser_country_or_area", "primary_industry_name", "question_number",
            "row_order", "row_name", "public_status", "column_header", "content_full",
        ]
        for question, record_type in PARQUET_QUESTIONS.items():
            row_filter = (
                (ds.field("question_number") == question)
                & (ds.field("public_status") == "Public")
            )
            if record_type != "initiative":
                row_filter = row_filter & (ds.field("row_name") == "Climate change")
            scanner = dataset.scanner(
                columns=columns,
                filter=row_filter,
                batch_size=40_000,
            )
            groups: dict[tuple[str, int], dict] = {}
            headers_seen: set[str] = set()
            for batch in scanner.to_batches():
                for cell in batch.to_pylist():
                    account = clean(cell["cdp_disclosing_org_number"])
                    if not account:
                        continue
                    row_order = int(cell["row_order"] or 1)
                    key = (account, row_order)
                    group = groups.setdefault(key, {"metadata": cell, "answers": {}, "headers": {}})
                    header = clean(cell["column_header"])
                    match = re.match(r"^(col\d+)_(.+)$", header)
                    if not match:
                        continue
                    field, label = match.groups()
                    headers_seen.add(header)
                    group["headers"][field] = label
                    value = clean(cell["content_full"])
                    if value:
                        values = group["answers"].setdefault(field, [])
                        if value not in values:
                            values.append(value)
            for header in sorted(headers_seen):
                audit.append({
                    "year": 2024, "question_id": question,
                    "source_file": str(path.relative_to(ROOT)),
                    "source_header": header,
                    "normalized_label": normalize_label(header.split("_", 1)[1]),
                })
            count = 0
            for (account, row_order), group in groups.items():
                if max_rows_per_question is not None and count >= max_rows_per_question:
                    break
                field_names = sorted(group["headers"], key=lambda name: int(name[3:]))
                row = tuple(" | ".join(group["answers"].get(field, [])) for field in field_names)
                label_map = {normalize_label(group["headers"][field]): i for i, field in enumerate(field_names)}
                metadata = group["metadata"]
                entry = {
                    "year": 2024,
                    "year_definition": "CDP questionnaire/disclosure year; fiscal reporting year may differ",
                    "cdp_account_number": account,
                    "organization": clean(metadata["disclosing_organization"]),
                    "country": clean(metadata["discloser_country_or_area"]),
                    "primary_industry": clean(metadata["primary_industry_name"]),
                    "question_id": question,
                    "record_type": record_type,
                    "question_row": row_order,
                    "source_file": str(path.relative_to(ROOT)),
                    "source_sheet": question,
                    "guide_file": str(GUIDES[2024][question].relative_to(ROOT)),
                    **currency_fields(currencies, 2024, account),
                }
                if record_type == "initiative":
                    initiative(row, label_map, entry)
                else:
                    risk_or_opportunity_2024(row, label_map, entry)
                if not entry["narrative_for_clustering"].strip():
                    continue
                entry["record_id"] = f"2024:{question}:{source_name}:{account}:{row_order}"
                write_entry(writer, complete_record(entry))
                count += 1
                total += 1
            print(f"2024 {source_name} {question}: {count:,} records", flush=True)
    return total


def extract(max_rows_per_sheet: int | None = None, output: Path = RECORDS) -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    glossary().to_csv(OUTPUT / "question_glossary.csv", index=False)
    currencies = load_reporting_currencies()
    audit: list[dict] = []
    total = 0
    partial_output = output.with_name(output.name + ".partial")
    with gzip.open(partial_output, "wt", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()
        for year, path in INPUTS.items():
            if not path.exists():
                raise FileNotFoundError(path)
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                sheets = LEGACY_SHEETS if year <= 2017 else MODERN_SHEETS
                for sheet in sheets:
                    if sheet not in book.sheetnames:
                        raise KeyError(f"Missing {year} question sheet {sheet}")
                    rows = book[sheet].iter_rows(values_only=True)
                    for header_row_number, headers in enumerate(rows, start=1):
                        if clean(headers[0]) in {"Account number", "program_name"}:
                            break
                    else:
                        raise ValueError(f"Header row not found: {year} {sheet}")
                    columns = field_headers(headers)
                    audit.extend({"year": year, "question_id": sheet, "source_file": str(path.relative_to(ROOT)), "source_header": clean(h), "normalized_label": normalize_label(clean(h).split(" - ", 2)[-1])} for h in headers if h)
                    count = 0
                    for excel_row, row in enumerate(rows, start=header_row_number + 1):
                        if max_rows_per_sheet is not None and count >= max_rows_per_sheet:
                            break
                        entry = common(row, columns, year, sheet, excel_row, currencies)
                        if not entry["cdp_account_number"]:
                            continue
                        if entry["record_type"] == "initiative":
                            initiative(row, columns, entry)
                        else:
                            risk_or_opportunity(row, columns, entry)
                        if not entry["narrative_for_clustering"].strip():
                            continue
                        entry["record_id"] = f"{year}:{sheet}:{excel_row}"
                        write_entry(writer, complete_record(entry))
                        count += 1
                        total += 1
                    print(f"{year} {sheet}: {count:,} records", flush=True)
            finally:
                book.close()
        total += extract_2024(writer, audit, max_rows_per_sheet, currencies)
    partial_output.replace(output)
    pd.DataFrame(audit).to_csv(OUTPUT / "source_field_mapping_audit.csv", index=False)
    write_coverage_report(output)
    print(f"Saved {total:,} rows to {output}", flush=True)


def write_coverage_report(input_path: Path = RECORDS) -> None:
    coverage = pd.read_csv(
        input_path,
        usecols=[
            "year", "record_type", "cdp_account_number",
            "estimated_annual_emissions_impact_tco2e", "emissions_boundary_reported",
            "cost_reported", "cost_text_reported", "timeframe_reported",
            "operational_complexity_signal_inferred",
            "condition_for_success_evidence_inferred",
        ],
        low_memory=False,
    )
    coverage = coverage.groupby(["year", "record_type"], as_index=False).agg(
        records=("cdp_account_number", "size"),
        companies=("cdp_account_number", "nunique"),
        estimated_annual_co2e_count=("estimated_annual_emissions_impact_tco2e", "count"),
        emissions_boundary_count=("emissions_boundary_reported", "count"),
        numeric_cost_count=("cost_reported", "count"),
        cost_text_count=("cost_text_reported", "count"),
        timeframe_count=("timeframe_reported", "count"),
        explicit_complexity_signal_count=("operational_complexity_signal_inferred", "count"),
        condition_evidence_count=("condition_for_success_evidence_inferred", "count"),
    )
    coverage.to_csv(OUTPUT / "coverage_by_year.csv", index=False)


def enrich_existing_with_currency() -> None:
    """Add explicit initiative comments and company-year currency without reclustering.

    The current extracted narrative already contains every nonblank initiative
    comment. Preserve those fitted topic assignments and leave the original
    processed files untouched; write suffixed enriched copies atomically.
    """
    currencies = load_reporting_currencies()
    if currencies.unavailable_years:
        print(f"Currency sources unavailable: {currencies.unavailable_years}", flush=True)
    extra_fields = [
        "initiative_comment_reported", "reporting_currency_raw",
        "reporting_currency_iso", "reporting_currency_source_file",
        "reporting_currency_source_sheet", "reporting_currency_source_question",
        "reporting_currency_status",
    ]
    sources = [
        RECORDS,
        OUTPUT / "climate_action_risk_opportunity_minilm_hybrid_clusters.csv.gz",
    ]
    coverage = defaultdict(Counter)
    for source_path in sources:
        if not source_path.exists():
            raise FileNotFoundError(source_path)
        destination = source_path.with_name(
            source_path.name.replace(".csv.gz", "_with_currency.csv.gz")
        )
        temporary = destination.with_name(destination.name + ".partial")
        with gzip.open(source_path, "rt", encoding="utf-8", newline="") as source, \
             gzip.open(temporary, "wt", encoding="utf-8", newline="") as target:
            reader = csv.DictReader(source)
            fieldnames = list(reader.fieldnames or [])
            if not {"year", "cdp_account_number", "record_type",
                    "reported_outcome_text", "narrative_for_clustering"} <= set(fieldnames):
                raise ValueError(f"Required columns are missing from {source_path}")
            writer = csv.DictWriter(
                target,
                fieldnames=fieldnames + [name for name in extra_fields if name not in fieldnames],
            )
            writer.writeheader()
            for row in reader:
                year = int(row["year"])
                if row["record_type"] == "initiative":
                    comment = row.get("initiative_comment_reported") or row["reported_outcome_text"]
                    if comment and comment not in row["narrative_for_clustering"]:
                        raise ValueError(
                            f"Initiative comment absent from clustering text: {row['record_id']}"
                        )
                else:
                    comment = ""
                row["initiative_comment_reported"] = comment
                row.update(currency_fields(currencies, year, row["cdp_account_number"]))
                writer.writerow(row)
                if source_path == RECORDS:
                    counts = coverage[year]
                    counts["records"] += 1
                    counts[f"currency_{row['reporting_currency_status']}"] += 1
                    if comment:
                        counts["initiative_comment_present"] += 1
        temporary.replace(destination)
        print(f"Saved {destination}", flush=True)

    summary = pd.DataFrame(
        [{"year": year, **counts} for year, counts in sorted(coverage.items())]
    ).fillna(0)
    summary["currency_coverage_pct"] = (
        100 * summary.get("currency_matched", 0) / summary["records"]
    ).round(2)
    summary.to_csv(OUTPUT / "reporting_currency_coverage_by_year.csv", index=False)
    print(summary.to_string(index=False), flush=True)


def cluster(input_path: Path = RECORDS, min_topic_size: int = 35, fit_limit: int = 6000) -> None:
    """Fit separate BERTopic models; transform every eligible record."""
    print("Loading BERTopic dependencies...", flush=True)
    import torch  # Load first on Windows, before sentence-transformers/UMAP.
    from bertopic import BERTopic
    from hdbscan import HDBSCAN
    from sentence_transformers import SentenceTransformer
    from sklearn.feature_extraction.text import CountVectorizer
    from umap import UMAP
    from src.cdp_text_clustering.umap_sklearn_compat import patch_topic_model_check_array

    patch_topic_model_check_array()

    del torch  # Import order is intentional.
    print("Reading extracted records...", flush=True)
    frame = pd.read_csv(
        input_path,
        usecols=["record_id", "record_type", "narrative_for_clustering"],
        dtype="string", low_memory=False,
    )
    assignments: dict[str, tuple[int, str]] = {}
    print("Loading multilingual embedding model...", flush=True)
    embedder = SentenceTransformer(MODEL_NAME)
    embedder.max_seq_length = 256
    print("Embedding model ready.", flush=True)
    topic_rows = []
    for record_type, subset in frame.groupby("record_type", sort=False):
        eligible = subset["narrative_for_clustering"].fillna("").str.len().ge(35)
        indices = subset.index[eligible]
        if len(indices) < max(min_topic_size * 2, 20):
            print(f"Skipping {record_type}: only {len(indices)} usable narratives")
            continue
        documents = frame.loc[indices, "narrative_for_clustering"].astype(str).tolist()
        print(f"Encoding {record_type}: {len(documents):,} records...", flush=True)
        embeddings = embedder.encode(documents, batch_size=32, normalize_embeddings=True, show_progress_bar=True)
        rng = np.random.default_rng(42)
        fit_indices = np.sort(rng.choice(len(documents), min(fit_limit, len(documents)), replace=False))
        model = BERTopic(
            embedding_model=embedder,
            umap_model=UMAP(n_neighbors=15, n_components=5, min_dist=0, metric="cosine", random_state=42, low_memory=True),
            hdbscan_model=HDBSCAN(min_cluster_size=min_topic_size, min_samples=10, prediction_data=True),
            vectorizer_model=CountVectorizer(stop_words="english", ngram_range=(1, 2), min_df=3),
            calculate_probabilities=False,
            verbose=False,
        )
        print(f"Fitting {record_type} topics on {len(fit_indices):,} records...", flush=True)
        model.fit_transform([documents[i] for i in fit_indices], embeddings=embeddings[fit_indices])
        print(f"Assigning {record_type} topics to all records...", flush=True)
        topics, _ = model.transform(documents, embeddings=embeddings)
        info = model.get_topic_info()
        names = info.set_index("Topic")["Name"].to_dict()
        for record_id, topic in zip(frame.loc[indices, "record_id"].astype(str), topics):
            assignments[record_id] = (int(topic), names.get(topic, "Outlier / unassigned"))
        info.insert(0, "record_type", record_type)
        topic_rows.append(info)
        model.save(str(OUTPUT / f"bertopic_{record_type}"), serialization="safetensors", save_ctfidf=True, save_embedding_model=MODEL_NAME)
        print(f"{record_type}: assigned {len(indices):,} records to {sum(info['Topic'].ge(0))} topics", flush=True)
    destination = OUTPUT / "climate_action_risk_opportunity_clustered.csv.gz"
    partial = destination.with_name(destination.name + ".partial")
    with gzip.open(input_path, "rt", encoding="utf-8", newline="") as source, gzip.open(partial, "wt", encoding="utf-8", newline="") as target:
        reader = csv.DictReader(source)
        writer = csv.DictWriter(target, fieldnames=list(reader.fieldnames or []) + ["topic_id", "topic_label"])
        writer.writeheader()
        for row in reader:
            topic_id, topic_label = assignments.get(row["record_id"], ("", ""))
            row.update({"topic_id": topic_id, "topic_label": topic_label})
            writer.writerow(row)
    partial.replace(destination)
    if topic_rows:
        pd.concat(topic_rows, ignore_index=True).to_csv(OUTPUT / "topic_summary.csv", index=False)
    print(f"Saved clustered dataset: {destination}")


def baseline_cluster(input_path: Path = RECORDS, clusters_per_type: int = 14, fit_limit: int = 30_000) -> None:
    """Memory-safe, fully local TF-IDF clustering when BERTopic cannot start.

    These are lexical clusters, not BERTopic topics or causal categories.
    """
    os.environ.setdefault("LOKY_MAX_CPU_COUNT", "8")
    from sklearn.cluster import MiniBatchKMeans
    from sklearn.feature_extraction.text import TfidfVectorizer

    frame = pd.read_csv(
        input_path,
        usecols=["record_id", "record_type", "narrative_for_clustering"],
        dtype="string",
        low_memory=False,
    )
    assignments: dict[str, tuple[str, str]] = {}
    summaries = []
    for record_type, subset in frame.groupby("record_type", sort=False):
        indices = subset.index[subset["narrative_for_clustering"].fillna("").str.len().ge(20)]
        if len(indices) < clusters_per_type * 3:
            continue
        documents = frame.loc[indices, "narrative_for_clustering"].astype(str).tolist()
        rng = np.random.default_rng(42)
        fit_positions = np.sort(rng.choice(len(documents), min(fit_limit, len(documents)), replace=False))
        training = [documents[i] for i in fit_positions]
        vectorizer = TfidfVectorizer(
            lowercase=True, strip_accents="unicode", stop_words="english",
            ngram_range=(1, 2), min_df=5, max_df=0.8,
            max_features=18000, sublinear_tf=True, dtype=np.float32,
        )
        matrix = vectorizer.fit_transform(training)
        model = MiniBatchKMeans(n_clusters=clusters_per_type, batch_size=1024, n_init=5, random_state=42)
        model.fit(matrix)
        terms = np.asarray(vectorizer.get_feature_names_out())
        labels = {}
        for topic in range(clusters_per_type):
            term_ids = np.argsort(model.cluster_centers_[topic])[-5:][::-1]
            labels[topic] = ", ".join(terms[term_ids])
        counts = np.zeros(clusters_per_type, dtype=int)
        record_ids = frame.loc[indices, "record_id"].astype(str).tolist()
        for start in range(0, len(documents), 2000):
            batch_labels = model.predict(vectorizer.transform(documents[start:start + 2000]))
            for record_id, topic in zip(record_ids[start:start + 2000], batch_labels):
                topic = int(topic)
                assignments[record_id] = (f"{record_type}_{topic:02d}", labels[topic])
                counts[topic] += 1
        for topic in range(clusters_per_type):
            summaries.append({
                "record_type": record_type,
                "cluster_id": f"{record_type}_{topic:02d}",
                "cluster_label": labels[topic],
                "records": int(counts[topic]),
                "method": "TF-IDF + MiniBatchKMeans",
            })
        print(f"{record_type}: clustered {len(indices):,} records", flush=True)
    destination = OUTPUT / "climate_action_risk_opportunity_lexical_clusters.csv.gz"
    partial = destination.with_name(destination.name + ".partial")
    with gzip.open(input_path, "rt", encoding="utf-8", newline="") as source, gzip.open(partial, "wt", encoding="utf-8", newline="") as target:
        reader = csv.DictReader(source)
        writer = csv.DictWriter(target, fieldnames=list(reader.fieldnames or []) + ["cluster_id", "cluster_label", "cluster_method"])
        writer.writeheader()
        for row in reader:
            cluster_id, cluster_label = assignments.get(row["record_id"], ("", ""))
            row.update({"cluster_id": cluster_id, "cluster_label": cluster_label, "cluster_method": "TF-IDF + MiniBatchKMeans (lexical baseline)"})
            writer.writerow(row)
    partial.replace(destination)
    pd.DataFrame(summaries).to_csv(OUTPUT / "lexical_cluster_summary.csv", index=False)
    print(f"Saved lexical-cluster dataset: {destination}", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", choices=["extract", "coverage", "baseline", "cluster", "enrich", "all"], default="extract")
    parser.add_argument("--max-rows-per-sheet", type=int, help="Development smoke test; do not use for production dataset")
    parser.add_argument("--min-topic-size", type=int, default=35)
    parser.add_argument("--fit-limit", type=int, default=6000, help="BERTopic training narratives per type (baseline uses at least 30,000); all eligible records receive assignments")
    args = parser.parse_args()
    path = RECORDS if args.max_rows_per_sheet is None else OUTPUT / "trial_records.csv.gz"
    if args.stage in {"extract", "all"}:
        extract(args.max_rows_per_sheet, path)
    if args.stage == "coverage":
        write_coverage_report(path)
    if args.stage == "enrich":
        enrich_existing_with_currency()
    if args.stage == "baseline":
        baseline_cluster(path, fit_limit=max(args.fit_limit, 30_000))
    if args.stage in {"cluster", "all"}:
        cluster(path, args.min_topic_size, args.fit_limit)


if __name__ == "__main__":
    main()
