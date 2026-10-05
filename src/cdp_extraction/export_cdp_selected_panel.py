"""Restore explicitly retained answers, join CDP classifications, and export CSV.

Run from the project root with Python including openpyxl and pyarrow:
    python -m src.cdp_extraction.export_cdp_selected_panel
"""
import csv
import gzip
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
import pyarrow.parquet as pq

try:
    from .cdp_wide_answers import KEEP_COLUMNS, DROP_COLUMNS, answer_only, missing_value
except ImportError:
    from src.cdp_extraction.cdp_wide_answers import KEEP_COLUMNS, DROP_COLUMNS, answer_only, missing_value

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "data/processed/cdp_climate_company_panel"


def account(value):
    text = str(value or "").strip()
    return text[:-2] if text.endswith(".0") else text


def classifications(year, accounts):
    root = ROOT / "data/raw/CDP" / str(year)
    result = {}

    def add(identifier, sector, industry):
        key = account(identifier)
        if key not in accounts:
            return
        pair = (str(sector or "").strip(), str(industry or "").strip())
        if key in result and result[key] != pair:
            raise ValueError(f"Conflicting classification for {year}/{key}")
        result[key] = pair

    if year >= 2024:
        folder = root if year == 2024 else root / "Climate Change"
        paths = sorted(folder.glob("*c_isin*summary*.parquet"))
        if len(paths) != 1:
            raise ValueError(f"Expected one CDP summary for {year}: {paths}")
        for row in pq.read_table(paths[0], columns=["cdp_disclosing_org_number", "primary_industry_name"]).to_pylist():
            add(row["cdp_disclosing_org_number"], "", row["primary_industry_name"])
        return result
    paths = sorted(root.glob("*.xlsx"))
    if len(paths) != 1:
        # Prefer the named final revisions used by the answer extraction.
        preferred = {2022: "CDP_2022_ClimateChange_Public_Inv_SC_v2.2.xlsx",
                     2023: "CDP_2023_ClimateChange_Public_Inv_SCv3.xlsx"}
        paths = [root / preferred[year]] if year in preferred else paths
    if len(paths) != 1:
        raise ValueError(f"Ambiguous workbook for {year}: {paths}")
    workbook = load_workbook(paths[0], read_only=True, data_only=True)
    try:
        sheet = workbook["Summary Data" if "Summary Data" in workbook.sheetnames else "Summary"]
        iterator = sheet.iter_rows(values_only=True)
        indices = None
        for _ in range(4):
            row = next(iterator)
            labels = [str(v or "").strip().lower() for v in row]
            account_label = next((c for c in ("account number", "account_id") if c in labels), None)
            if account_label:
                indices = (labels.index(account_label),
                           labels.index("primary sector") if "primary sector" in labels else None,
                           labels.index("primary industry") if "primary industry" in labels else None)
                print(f"{year}: sector available={indices[1] is not None}, industry available={indices[2] is not None}", flush=True)
                break
        if indices is None:
            raise ValueError(f"Account header missing in {paths[0]}")
        a, s, i = indices
        for row in iterator:
            add(row[a], row[s] if s is not None else "", row[i] if i is not None else "")
    finally:
        workbook.close()
    return result


def export_panel(output_dir=OUTPUT):
    csv.field_size_limit(100_000_000)
    panel_path = output_dir / "cdp_trucost_factset_lseg_climate_2016_2025.csv.gz"
    csv_path = panel_path.with_suffix("")
    with gzip.open(panel_path, "rt", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        original_columns = reader.fieldnames
        rows = list(reader)
    mapping_path = output_dir / "cdp_column_name_mapping.csv"
    with mapping_path.open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        mapping_fields = reader.fieldnames
        mapping = list(reader)
    lookup = {(r['family'], r['question_code'], r['column_header']): r['column_name']
              for r in mapping if r['column_name'] in KEEP_COLUMNS}
    assert set(lookup.values()) == set(KEEP_COLUMNS), set(KEEP_COLUMNS) - set(lookup.values())
    answers = defaultdict(lambda: defaultdict(list))
    with gzip.open(output_dir / "cdp_climate_question_answers_2016_2025.csv.gz", "rt", encoding="utf-8", newline="") as stream:
        for r in csv.DictReader(stream):
            col = lookup.get((r['family'], r['question_code'], r['column_header']))
            if col:
                answers[r['year'], account(r['cdp_account_number'])][col].append(r['answer'])
    columns = [c for c in original_columns if c not in DROP_COLUMNS]
    columns += [c for c in KEEP_COLUMNS if c not in columns]
    columns += [c for c in ('primary_sector', 'primary_industry') if c not in columns]
    by_year = defaultdict(list)
    for row in rows:
        by_year[int(row['year'])].append(row)
    metadata = {}
    for year, year_rows in sorted(by_year.items()):
        data = classifications(year, {account(r['cdp_account_number']) for r in year_rows})
        metadata.update({(str(year), k): v for k, v in data.items()})
        print(f"{year}: classified {len(data):,} CDP accounts", flush=True)
    pending = csv_path.with_name(csv_path.name + '.partial')
    coverage = defaultdict(lambda: [0, 0, 0])
    with pending.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            key = row['year'], account(row['cdp_account_number'])
            result = {c: answer_only(row.get(c, '')) if c.startswith('cdp_') else row.get(c, '') for c in columns}
            for col in KEEP_COLUMNS:
                values = answers[key].get(col, [])
                result[col] = (values[0] if len(values) == 1 else
                               json.dumps(values, ensure_ascii=False, separators=(',', ':')) if values else '')
            result['primary_sector'], result['primary_industry'] = metadata.get(key, ('', ''))
            writer.writerow(result)
            coverage[row['year']][0] += 1
            coverage[row['year']][1] += bool(result['primary_sector'])
            coverage[row['year']][2] += bool(result['primary_industry'])
    # Verify all retained source values, original row order, and untouched fields.
    missing = defaultdict(int)
    with pending.open(encoding='utf-8', newline='') as stream:
        reader = csv.DictReader(stream)
        checked = 0
        for before, after in zip(rows, reader, strict=True):
            key = before['year'], account(before['cdp_account_number'])
            for col in columns:
                missing[col] += missing_value(after[col])
                if col in KEEP_COLUMNS:
                    values = answers[key].get(col, [])
                    actual = json.loads(after[col]) if len(values) > 1 else [after[col]] if values else []
                    assert actual == values, (key, col)
                elif col in original_columns and col not in {'primary_sector', 'primary_industry'}:
                    assert after[col] == answer_only(before[col]) if col.startswith('cdp_') else after[col] == before[col]
            checked += 1
    assert checked == len(rows) and not (set(columns) & DROP_COLUMNS)
    pending.replace(csv_path)
    # Keep compressed and uncompressed versions identical.
    compressed = panel_path.with_name(panel_path.name + '.partial')
    with csv_path.open('rb') as source, gzip.open(compressed, 'wb') as target:
        import shutil
        shutil.copyfileobj(source, target)
    compressed.replace(panel_path)
    with mapping_path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=mapping_fields)
        writer.writeheader()
        for row in mapping:
            row['retained'] = row['column_name'] in columns
            writer.writerow(row)
    with (output_dir / 'cdp_selected_panel_missingness.csv').open('w', encoding='utf-8', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['column_name', 'missing_rows', 'total_rows', 'missing_pct', 'explicit_keep'])
        for col in columns:
            writer.writerow([col, missing[col], len(rows), 100 * missing[col] / len(rows), col in KEEP_COLUMNS])
    print('Classification coverage (year, rows, sector, industry):', dict(coverage), flush=True)
    print(f'Validated and saved {len(rows):,} rows, {len(columns)} columns: {csv_path}', flush=True)


if __name__ == '__main__':
    export_panel()
