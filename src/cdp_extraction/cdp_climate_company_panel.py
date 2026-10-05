"""Export a traceable 2016--2025 CDP/Trucost/FactSet/LSEG climate panel.

The company-year file includes question-specific CDP answer columns. Every
selected answer is also retained in the long file with its original labels.
Questionnaire families change in 2024, so a missing family is not coded No.
"""

from __future__ import annotations

import csv
import gc
import gzip
import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
import pyarrow.dataset as ds
from openpyxl import load_workbook

from src.cdp_extraction.extract_cdp_climate_actions_2016_2024 import INPUTS


YEARS = tuple(range(2016, 2026))
FAMILIES = ("targets", "emissions", "verification", "carbon_regulation",
            "carbon_credits", "internal_carbon_price", "engagement")
LONG_COLUMNS = ("year", "cdp_account_number", "question_code", "family",
                "row_order", "row_name", "column_header", "answer", "source_file")
OUTPUT_DIR = Path(__file__).resolve().parents[2] / "data/processed/cdp_climate_company_panel"


def account_key(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    value = str(value).strip()
    return value[:-2] if value.endswith(".0") else value


def family_for(year: int, code: str) -> str | None:
    code = str(code).upper().removeprefix("Q")
    if year <= 2017:
        if code.startswith("CC3.1"):
            return "targets"
        if code == "CC8.3A":
            return "emissions"
        if code.startswith(("CC8.6", "CC8.7", "CC8.8")):
            return "verification"
        if code.startswith("CC13.1"):
            return "carbon_regulation"
        if code.startswith("CC13.2"):
            return "carbon_credits"
        return None
    if year <= 2023:
        if re.fullmatch(r"C4\.1[A-C]?", code):
            return "targets"
        if code == "C6.3":
            return "emissions"
        if code.startswith("C10.1"):
            return "verification"
        if code in {"C11.1", "C11.1A"}:
            return "carbon_regulation"
        if code in {"C11.2", "C11.2A"}:
            return "carbon_credits"
        if code in {"C11.3", "C11.3A"}:
            return "internal_carbon_price"
        if re.fullmatch(r"C12\.1[A-D]?", code):
            return "engagement"
        return None
    if code.startswith(("7.53", "20.16")):
        return "targets"
    if code in {"7.7", "20.5"}:
        return "emissions"
    if code.startswith("7.9") or code == "20.8":
        return "verification"
    if code in {"3.5", "3.5.1"}:
        return "carbon_regulation"
    if code.startswith("7.79"):
        return "carbon_credits"
    if code in {"5.10", "5.10.1"}:
        return "internal_carbon_price"
    if code.startswith("5.11") or code == "18.3":
        return "engagement"
    return None


def selected_legacy_sheets(year: int, sheetnames: list[str]) -> list[str]:
    # Include only actual questionnaire code sheets, not module intro sheets.
    return [name for name in sheetnames
            if re.fullmatch(r"C{1,2}\d+(?:\.\d+)*[a-z]?", name, re.I)
            and family_for(year, name)]


def parquet_sources(year: int, root: Path) -> list[Path]:
    if year == 2024:
        sources = sorted((root / "2024").glob("*c_isin*responses*.parquet"))
    else:
        sources = []
        for folder, prefix in (("Climate Change", "*c_isin*responses*.parquet"),
                               ("00 Integrated Questions", "*gen_isin*responses*.parquet")):
            sources.extend(sorted((root / "2025" / folder).glob(prefix)))
    if not sources:
        raise FileNotFoundError(f"No CDP response Parquet for {year} under {root}")
    return sources


def legacy_answers(year: int, path: Path, accounts: set[str]):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for code in selected_legacy_sheets(year, workbook.sheetnames):
            sheet = workbook[code]
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            header_row = 1 if any(str(v).strip().casefold() in
                                  {"account_id", "account number"} for v in first) else 2
            headers = [str(v or "") for v in next(sheet.iter_rows(
                min_row=header_row, max_row=header_row, values_only=True))]
            account_col = next((i for i, h in enumerate(headers) if h.strip().casefold()
                                in {"account_id", "account number"}), None)
            if account_col is None:
                continue
            row_col = next((i for i, h in enumerate(headers) if h.strip().casefold()
                            == "row"), None)
            row_name_col = next((i for i, h in enumerate(headers) if h.strip().casefold()
                                 == "rowname"), None)
            family = family_for(year, code)
            answer_cols = [(i, h) for i, h in enumerate(headers)
                           if h.upper().startswith((code.upper() + "_", code.upper() + " "))]
            if not answer_cols:
                # Some single-answer sheets use a bare question label.
                answer_cols = [(i, h) for i, h in enumerate(headers)
                               if i >= 6 and h.upper().startswith(code.upper())]
            for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                account = account_key(values[account_col])
                if account not in accounts:
                    continue
                row_order = values[row_col] if row_col is not None else ""
                row_name = values[row_name_col] if row_name_col is not None else ""
                for i, header in answer_cols:
                    if i >= len(values) or values[i] is None or str(values[i]).strip() == "":
                        continue
                    yield (year, account, code, family, row_order, row_name,
                           header, str(values[i]).strip(), str(path))
    finally:
        workbook.close()


def recent_answers(year: int, root: Path, accounts: set[str]):
    question_codes = (
        ["Q7.53", "Q7.53.1", "Q7.53.2", "Q7.53.3", "Q7.53.4",
         "Q20.16", "Q20.16.1", "Q20.16.2", "Q20.16.3", "Q7.7", "Q20.5",
         "Q7.9", "Q7.9.1", "Q7.9.2", "Q7.9.3", "Q20.8",
         "Q3.5", "Q3.5.1", "Q7.79", "Q7.79.1", "Q5.10", "Q5.10.1",
         "Q18.3"] + [f"Q5.11.{i}" for i in range(1, 10)] + ["Q5.11"]
    )
    for path in parquet_sources(year, root):
        dataset = ds.dataset(path)
        columns = ["cdp_disclosing_org_number", "question_number", "row_order",
                   "row_name", "column_header", "content_full"]
        # A filter on matched CDP accounts is pushed into the Arrow scan.
        numeric_accounts = [int(account) for account in accounts if account.isdigit()]
        filt = (ds.field("cdp_disclosing_org_number").isin(numeric_accounts)
                & ds.field("question_number").isin(question_codes))
        for batch in dataset.scanner(columns=columns, filter=filt,
                                     batch_size=10_000, use_threads=False).to_batches():
            for item in batch.to_pylist():
                code = item["question_number"] or ""
                family = family_for(year, code)
                answer = item["content_full"]
                if family is None or answer is None or str(answer).strip() == "":
                    continue
                yield (year, account_key(item["cdp_disclosing_org_number"]),
                       code, family, item["row_order"], item["row_name"],
                       item["column_header"], str(answer).strip(), str(path))


def build_climate_panel(cdp_root: Path, cache_path: Path, output_dir: Path = OUTPUT_DIR):
    try:
        from .common_companies_by_year import build_cdp_trucost_factset_common_panel
    except ImportError:
        from src.dataset_readers.common_companies_by_year import build_cdp_trucost_factset_common_panel

    output_dir.mkdir(parents=True, exist_ok=True)
    match_path = (output_dir / "matching" /
                  "cdp_trucost_factset_common_2016_2025.csv.gz")
    if match_path.exists():
        common = pd.read_csv(match_path, low_memory=False)
        if "cdp_account_number" not in common.columns:
            raise KeyError(f"Saved match lacks CDP account IDs: {match_path}")
        print(f"Reusing {len(common):,} saved CDP company-year matches", flush=True)
    else:
        common, _ = build_cdp_trucost_factset_common_panel(
            years=YEARS, cdp_root=cdp_root, matched_panel_cache_path=cache_path,
            output_dir=output_dir / "matching", fuzzy_threshold=95.0,
        )
    if common.empty:
        raise ValueError("No CDP/Trucost/FactSet company-years matched")
    common["cdp_account_number"] = common["cdp_account_number"].map(account_key)
    # The cache already holds the LSEG match, ISIN, ESG and Trucost emissions.
    required = ["year", "cdp_account_number", "cdp_org_name", "cdp_match_method",
                "cdp_name_score", "trucost_company_id", "trucost_name",
                "factset_entity_id", "factset_name", "isin", "lseg_name", "esg_score",
                "scope_1_emissions", "scope_2_emissions", "scope_3_upstream_emissions",
                "scope_3_downstream_emissions"]
    missing = [col for col in required if col not in common.columns]
    if missing:
        raise KeyError(f"Matched cache lacks required fields: {missing}")
    output = common[required + [c for c in
        ("lseg_match_method", "lseg_name_score", "factset_match_method",
         "factset_name_score", "scope_2_location_based_emissions",
         "scope_2_market_based_emissions") if c in common.columns]].copy()
    output = output.rename(columns={"isin": "lseg_isin"})
    output["cdp_match_score"] = pd.to_numeric(output.pop("cdp_name_score"), errors="coerce")
    output["cdp_response_rows"] = 0
    for family in FAMILIES:
        output[f"cdp_{family}_answer_count"] = 0
        output[f"cdp_{family}_question_count"] = 0

    answer_count = Counter()
    question_sets = defaultdict(set)
    long_path = output_dir / "cdp_climate_question_answers_2016_2025.csv.gz"
    pending_long_path = long_path.with_name(long_path.name + ".partial")
    with gzip.open(pending_long_path, "wt", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(LONG_COLUMNS)
        for year in YEARS:
            accounts = set(output.loc[output["year"].eq(year), "cdp_account_number"]) - {""}
            if not accounts:
                continue
            if year <= 2023:
                if not INPUTS[year].exists():
                    raise FileNotFoundError(INPUTS[year])
                answers = legacy_answers(year, INPUTS[year], accounts)
            else:
                answers = recent_answers(year, cdp_root, accounts)
            year_count = 0
            for answer in answers:
                writer.writerow(answer)
                key = (year, answer[1], answer[3])
                answer_count[key] += 1
                question_sets[key].add(answer[2])
                year_count += 1
            print(f"{year}: {len(accounts):,} matched CDP accounts; {year_count:,} selected answers", flush=True)
            del answers
            gc.collect()
    pending_long_path.replace(long_path)

    for family in FAMILIES:
        keys = list(zip(output["year"], output["cdp_account_number"],
                        [family] * len(output)))
        output[f"cdp_{family}_answer_count"] = [answer_count[key] for key in keys]
        output[f"cdp_{family}_question_count"] = [len(question_sets[key]) for key in keys]
    output["cdp_response_rows"] = output[[f"cdp_{f}_answer_count" for f in FAMILIES]].sum(axis=1)

    panel_path = output_dir / "cdp_trucost_factset_lseg_climate_2016_2025.csv.gz"
    output.to_csv(panel_path, index=False, compression="gzip")
    summary = output.groupby("year", as_index=False).agg(
        company_years=("cdp_account_number", "size"),
        cdp_accounts=("cdp_account_number", "nunique"),
        with_lseg_esg=("esg_score", "count"),
        with_cdp_selected_answers=("cdp_response_rows", lambda x: int(x.gt(0).sum())),
    )
    summary_path = output_dir / "climate_panel_coverage_by_year.csv"
    summary.to_csv(summary_path, index=False)
    print(summary.to_string(index=False))
    print(f"Saved panel: {panel_path}\nSaved CDP answers: {long_path}\nSaved coverage: {summary_path}")
    return refresh_harmonized_panel(output_dir)


def refresh_harmonized_panel(output_dir: Path = OUTPUT_DIR) -> pd.DataFrame:
    """Replace compact CDP fields using the saved original-answer export."""
    try:
        from .cdp_wide_answers import expand_panel
    except ImportError:
        from src.cdp_extraction.cdp_wide_answers import expand_panel
    panel_path = output_dir / "cdp_trucost_factset_lseg_climate_2016_2025.csv.gz"
    long_path = output_dir / "cdp_climate_question_answers_2016_2025.csv.gz"
    expand_panel(panel_path, long_path)
    try:
        from .export_cdp_selected_panel import export_panel
    except ImportError:
        from src.cdp_extraction.export_cdp_selected_panel import export_panel
    export_panel(output_dir)
    return pd.read_csv(panel_path, low_memory=False)
