"""Stage matched-company CDP answers for four curated section datasets."""
import csv
import gzip
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'data/processed/cdp_section_datasets'
STAGE = OUT / '_staging'


def key(value):
    text = str(value or '').strip()
    return text[:-2] if text.endswith('.0') else text


def clean(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value).strip()


def section(code):
    c = code.upper().removeprefix('Q')
    if c.startswith(('CC3.', 'C4.1', 'C4.2', 'C4.3', '7.53', '7.54', '7.55', '20.16', '20.17')):
        return 'targets_performance'
    if c.startswith(('CC8.6', 'CC8.7', 'CC8.8', 'CC14.2', 'C10.', '7.9')) or c == '20.8':
        return 'verification'
    if c.startswith(('CC13.', 'C11.', '3.5', '7.79', '5.10')) or c in {'CC2.2C', 'CC2.2D'}:
        return 'carbon_pricing'
    if c.startswith(('CC14.4', 'C12.1', 'C12.2', '5.11', '18.3')):
        return 'engagement'
    return None


def base_sample():
    path = ROOT / 'data/processed/cdp_climate_company_panel/cdp_trucost_factset_lseg_climate_2016_2025.csv.gz'
    with gzip.open(path, 'rt', encoding='utf-8', newline='') as stream:
        return list(csv.DictReader(stream))


def legacy(year, accounts, only_codes=None):
    paths = [p for p in (ROOT / f'data/raw/CDP/{year}').glob('*.xlsx') if not p.name.startswith('~$')]
    if year == 2022: paths = [p for p in paths if 'v2.2' in p.name]
    if year == 2023: paths = [p for p in paths if 'SCv3' in p.name]
    if len(paths) != 1: raise ValueError(paths)
    path = paths[0]
    workbook = load_workbook(path, read_only=True, data_only=True)
    currencies = {}
    count = 0
    suffix = '_supplement' if only_codes else ''
    pending = STAGE / f'{year}{suffix}.jsonl.gz.partial'
    with gzip.open(pending, 'wt', encoding='utf-8') as dest:
        for name in workbook.sheetnames:
            if only_codes and name.upper() not in only_codes: continue
            if not (name.startswith(('CC0', 'C0', 'CC2', 'CC3', 'CC5', 'CC6', 'CC8', 'CC13', 'CC14', 'C2', 'C4', 'C10', 'C11', 'C12')) or name in {'Summary', 'Summary Data'}):
                continue
            iterator = workbook[name].iter_rows(values_only=True)
            header = None
            for _ in range(3):
                values = next(iterator, None)
                if values is None: break
                labels = [str(v or '').strip() for v in values]
                matches = [i for i,h in enumerate(labels) if h.lower() in {'account number', 'account_id'}]
                if matches:
                    header, account_index = labels, matches[0]
                    break
            if header is None: continue
            low = [h.lower() for h in header]
            row_index = low.index('row') if 'row' in low else None
            row_name_index = low.index('rowname') if 'rowname' in low else None
            currency_cols = [i for i,h in enumerate(low) if h == 'currency' or ('select' in h and 'currency' in h)]
            selected = []
            for i,h in enumerate(header):
                m = re.match(r'^((?:CC|C)\d+(?:\.\d+)*[a-z]?)(?:_|\s|$)', h, re.I)
                if m and section(m[1]) and (not only_codes or m[1].upper() in only_codes): selected.append((i,m[1],h))
            if not selected and not currency_cols: continue
            for values in iterator:
                if account_index >= len(values): continue
                account = key(values[account_index])
                if account not in accounts: continue
                for i in currency_cols:
                    if i < len(values) and values[i] is not None and str(values[i]).strip():
                        currencies.setdefault(account, clean(values[i]))
                groups = defaultdict(dict)
                for i,code,h in selected:
                    if i < len(values) and values[i] is not None and str(values[i]).strip():
                        groups[code][h] = clean(values[i])
                for code,fields in groups.items():
                    row = dict(year=year, cdp_account_number=account, question_code=code,
                               row_order=clean(values[row_index]) if row_index is not None and values[row_index] is not None else '',
                               row_name=clean(values[row_name_index]) if row_name_index is not None and values[row_name_index] is not None else '',
                               source_file=str(path.relative_to(ROOT)), source_sheet=name, fields=fields)
                    dest.write(json.dumps(row,ensure_ascii=False)+'\n')
                    count += 1
            print(f'{year}: staged {name}; {count:,} records', flush=True)
    workbook.close()
    pending.replace(STAGE / f'{year}{suffix}.jsonl.gz')
    if not only_codes:
        (STAGE / f'{year}_currencies.json').write_text(json.dumps(currencies), encoding='utf-8')


def recent(year, accounts, prefixes=None):
    import pyarrow.dataset as ds
    root = ROOT / f'data/raw/CDP/{year}'
    paths = (list(root.glob('*c_isin*responses*.parquet')) if year == 2024 else
             list((root / 'Climate Change').glob('*c_isin*responses*.parquet')) +
             list((root / '00 Integrated Questions').glob('*gen_isin*responses*.parquet')))
    codes = []
    for prefix in (prefixes or ['7.53', '7.54', '7.55', '7.9', '3.5', '7.79', '5.10', '5.11', '20.16', '20.17', '20.8', '18.3']):
        codes += ['Q'+prefix] + ['Q'+prefix+'.'+str(i) for i in range(1,15)]
    pending = STAGE / f'{year}.jsonl.gz.partial'
    currencies = {}
    summaries = list(root.glob('*c_isin*summary*.parquet')) if year==2024 else list((root/'Climate Change').glob('*c_isin*summary*.parquet'))
    for path in summaries:
        for r in ds.dataset(path).to_table(columns=['cdp_disclosing_org_number','currency']).to_pylist():
            if key(r['cdp_disclosing_org_number']) in accounts and r['currency']:
                currencies[key(r['cdp_disclosing_org_number'])] = str(r['currency'])
    total=0
    with gzip.open(pending,'wt',encoding='utf-8') as dest:
        for path in paths:
            dataset=ds.dataset(path)
            columns=['cdp_disclosing_org_number','question_number','row_order','row_name','column_header','content_full']
            filt=ds.field('cdp_disclosing_org_number').isin([int(a) for a in accounts]) & ds.field('question_number').isin(codes)
            grouped=defaultdict(lambda: defaultdict(list))
            for batch in dataset.scanner(columns=columns,filter=filt,batch_size=10000,use_threads=False).to_batches():
                for r in batch.to_pylist():
                    if r['content_full'] is None or not str(r['content_full']).strip():continue
                    k=(key(r['cdp_disclosing_org_number']),r['question_number'],str(r['row_order'] or ''),str(r['row_name'] or ''))
                    grouped[k][str(r['column_header'] or '')].append(str(r['content_full']).strip())
            for (account,code,order,name), fields in grouped.items():
                row=dict(year=year,cdp_account_number=account,question_code=code,row_order=order,row_name=name,source_file=str(path.relative_to(ROOT)),source_sheet='',fields={h:vs[0] if len(vs)==1 else vs for h,vs in fields.items()})
                dest.write(json.dumps(row,ensure_ascii=False)+'\n');total+=1
            print(f'{year}: staged {path.name}; {total:,} records',flush=True)
    pending.replace(STAGE/f'{year}.jsonl.gz')
    (STAGE/f'{year}_currencies.json').write_text(json.dumps(currencies),encoding='utf-8')


if __name__ == '__main__':
    STAGE.mkdir(parents=True, exist_ok=True)
    sample=base_sample()
    for year in range(2016,2026):
        if (STAGE/f'{year}.jsonl.gz').exists() and (STAGE/f'{year}_currencies.json').exists():
            print(f'{year}: reusing staged records',flush=True);continue
        accounts={key(r['cdp_account_number']) for r in sample if int(r['year'])==year}
        if year<2024:legacy(year,accounts)
        else:recent(year,accounts)
