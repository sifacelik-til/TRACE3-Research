"""
Extract CDP questionnaire answers for an organization across years (2014-2024)
and align questions over time using normalized question keys.

Usage:
  .venv\Scripts\python src\extract_cdp_org_answers_2014_2024.py --org-name "Volkswagen AG"
  .venv\Scripts\python src\extract_cdp_org_answers_2014_2024.py --org-name "Hitachi" --start-year 2014 --end-year 2024
"""
from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd
from difflib import get_close_matches
from openpyxl import load_workbook


logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
LOGGER = logging.getLogger(__name__)


DEFAULT_CDP_ROOT = Path(r"data\raw\CDP")
DEFAULT_OUT_DIR = Path(r"data\outputs")
FACTSET_ENTITY_PATH = Path(__file__).resolve().parent.parent / "data" / "raw" / "FactSet" / "sym_entity_v1_full_12328" / "sym_entity.txt"
QUESTION_CODE_RE = re.compile(r"^(CC?\d+(?:\.\d+)*[a-z]?)", re.IGNORECASE)
LEGACY_COL_CODE_RE = re.compile(r"^(CC?\d+(?:\.\d+)*[a-z]?)_C\d+_", re.IGNORECASE)

try:
    from .cdp_theme_taxonomy import THEME_PATTERNS, QUESTION_CODE_THEMES
except ImportError:
    try:
        from data_engine.cdp_theme_taxonomy import (
            THEME_PATTERNS,
            QUESTION_CODE_THEMES,
        )
    except ModuleNotFoundError:
        from cdp_theme_taxonomy import THEME_PATTERNS, QUESTION_CODE_THEMES


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def slugify(value: str) -> str:
    out = re.sub(r"[^a-zA-Z0-9]+", "_", value.strip()).strip("_").lower()
    return out or "org"


def resolve_org_name_from_factset_id(factset_id: str, path: Path = FACTSET_ENTITY_PATH) -> str:
    if not path.exists():
        raise FileNotFoundError(f"FactSet entity file not found: {path}")

    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            df = pd.read_csv(path, sep="|", dtype=str, encoding=encoding, low_memory=False)
            break
        except UnicodeDecodeError:
            continue
    else:
        df = pd.read_csv(path, sep="|", dtype=str, encoding="utf-8", low_memory=False, engine="python")

    df.columns = [c.strip().strip('"') for c in df.columns]
    df["FACTSET_ENTITY_ID"] = df["FACTSET_ENTITY_ID"].fillna("").astype(str).str.strip()
    row = df[df["FACTSET_ENTITY_ID"].str.upper() == factset_id.strip().upper()]
    if row.empty:
        raise ValueError(f"FactSet ID not found in entity file: {factset_id}")

    org_name = row.iloc[0].get("ENTITY_PROPER_NAME", "")
    if is_missing(org_name):
        raise ValueError(f"No organization name found for FactSet ID: {factset_id}")
    return str(org_name).strip()


def parse_year_file(year_dir: Path) -> Optional[Path]:
    """Return the primary Climate Change response file for one CDP year.

    Newer CDP releases contain several disclosure themes and frequently contain
    duplicate delivery folders. Selecting the first ``responses`` Parquet
    recursively can therefore select Biodiversity or Integrated Questions.
    """
    if not year_dir.exists():
        return None

    if year_dir.name in {"2024", "2025"}:
        parquet_candidates = [
            path
            for path in year_dir.rglob("*responses*.parquet")
            if (
                "climate change" in str(path.parent).lower()
                or "_c_isin_" in path.name.lower()
            )
        ]
        if parquet_candidates:
            return sorted(parquet_candidates)[0]

    xlsx_candidates = [p for p in year_dir.rglob("*.xlsx") if "glossary" not in p.name.lower()]
    if not xlsx_candidates:
        return None

    priority = [p for p in xlsx_candidates if "public" in p.name.lower()]
    if priority:
        return sorted(priority)[0]
    return sorted(xlsx_candidates)[0]


def find_header_row(ws, max_scan_rows: int = 6) -> Optional[int]:
    for row_idx in range(1, max_scan_rows + 1):
        row_vals = [ws.cell(row=row_idx, column=col).value for col in range(1, min(80, ws.max_column) + 1)]
        row_norm = [normalize_text(v) for v in row_vals if v is not None]
        if not row_norm:
            continue
        joined = " | ".join(row_norm)
        if "organization" in joined or "organisation" in joined:
            return row_idx
        if "account number" in joined or "account no" in joined:
            return row_idx
    return 1


def get_col_idx(headers: Sequence[str], keywords: Sequence[str]) -> Optional[int]:
    for idx, name in enumerate(headers):
        low = normalize_text(name)
        if any(k in low for k in keywords):
            return idx
    return None


def is_missing(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    text = str(value).strip()
    if not text:
        return True
    # CDP uses N/A and "Not applicable" as valid selected responses. Keep
    # them in the output; only values representing a missing cell are omitted.
    if text.lower() in {"nan", "none", "null"}:
        return True
    return False


def extract_question_code(sheet_name: str, column_header: str) -> str:
    col = str(column_header or "").strip()
    sheet = str(sheet_name or "").strip()

    m = LEGACY_COL_CODE_RE.match(col)
    if m:
        return m.group(1).upper()

    m = QUESTION_CODE_RE.match(col)
    if m:
        return m.group(1).upper()

    m = QUESTION_CODE_RE.match(sheet)
    if m:
        return m.group(1).upper()

    return sheet.upper()


def canonical_question_key(code: str) -> str:
    key = (code or "").strip().upper()
    key = key.replace(" ", "")
    if key.startswith("CC"):
        key = "C" + key[2:]
    key = key.rstrip(".")
    return key


def parse_legacy_column_text(column_header: str) -> str:
    text = str(column_header or "").strip()
    text = re.sub(r"^CC?\d+(?:\.\d+)*[a-z]?_C\d+_", "", text, flags=re.IGNORECASE)
    return text.strip(" -_")


def categorize_theme(sheet_name: str, question_text: str, answer: object) -> List[str]:
    haystack = " ".join([normalize_text(sheet_name), normalize_text(question_text), normalize_text(answer)])
    code = canonical_question_key(extract_question_code(sheet_name=sheet_name, column_header=question_text))

    tags: List[str] = []
    for theme, patterns in QUESTION_CODE_THEMES:
        if any(re.search(pattern, code, flags=re.IGNORECASE) for pattern in patterns):
            tags.append(theme)

    for tag, patterns in THEME_PATTERNS:
        if any(re.search(pattern, haystack, flags=re.IGNORECASE) for pattern in patterns):
            tags.append(tag)

    if not tags and "supplier" in haystack and any(k in haystack for k in ("transition", "renewable", "clean electricity", "procurement")):
        tags.append("supply_chain_transition")

    if "other" in tags:
        tags = [t for t in tags if t != "other"]

    return tags


def primary_theme(tags: Sequence[str]) -> str:
    return tags[0] if tags else "other"


def match_org_rows(
    org_series: Iterable[object], org_name: str, fuzzy_cutoff: float = 0.86
) -> Tuple[List[int], List[str]]:
    normalized_target = normalize_text(org_name)
    org_values = ["" if v is None else str(v) for v in org_series]
    org_norm = [normalize_text(v) for v in org_values]

    matches = []
    for i, n in enumerate(org_norm):
        if not n:
            continue
        if n == normalized_target or normalized_target in n or n in normalized_target:
            matches.append(i)

    if matches:
        labels = sorted({org_values[i] for i in matches if org_values[i]})
        return matches, labels

    non_empty_unique = sorted({n for n in org_norm if n})
    closest = get_close_matches(normalized_target, non_empty_unique, n=5, cutoff=fuzzy_cutoff)
    if not closest:
        return [], []

    close_set = set(closest)
    fuzzy_matches = [i for i, n in enumerate(org_norm) if n in close_set]
    labels = sorted({org_values[i] for i in fuzzy_matches if org_values[i]})
    return fuzzy_matches, labels


def extract_from_legacy_xlsx(
    year: int, file_path: Path, org_name: str
) -> Tuple[List[Dict[str, object]], List[str]]:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    summary_sheet_name = "Summary Data" if "Summary Data" in sheet_names else ("Summary" if "Summary" in sheet_names else None)

    matched_org_labels: List[str] = []
    account_ids: set[str] = set()
    org_norm_target = normalize_text(org_name)

    if summary_sheet_name:
        ws = wb[summary_sheet_name]
        header_row = find_header_row(ws)
        headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
        headers_str = ["" if h is None else str(h) for h in headers]
        org_idx = get_col_idx(headers_str, ("organization", "organisation", "response organisation"))
        acct_idx = get_col_idx(headers_str, ("account number", "account no", "account"))

        if org_idx is not None:
            org_values = []
            acct_values = []
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                org_values.append(row[org_idx] if org_idx < len(row) else None)
                acct_values.append(row[acct_idx] if (acct_idx is not None and acct_idx < len(row)) else None)

            matched_idx, matched_org_labels = match_org_rows(org_values, org_name)
            for i in matched_idx:
                acct = acct_values[i]
                if acct is not None and str(acct).strip():
                    account_ids.add(str(acct).strip())

    records: List[Dict[str, object]] = []
    metadata_terms = (
        "account number",
        "account no",
        "organisation",
        "organization",
        "country",
        "access",
        "public",
        "samples",
        "response received date",
        "activities",
        "sectors",
        "industries",
        "primary activity",
        "primary sector",
        "primary industry",
        "primary questionnaire sector",
        "ticker",
        "isin",
        "row",
        "rowname",
        "region",
        "status",
        "submitted date",
        "invitation status",
        "request response status",
        "attachments",
        "complexity",
        "version",
        "back to index",
    )

    for sheet_name in sheet_names:
        low_sheet = sheet_name.lower()
        if low_sheet in {"criteria", "response language"}:
            continue
        if low_sheet.startswith("summary"):
            continue

        ws = wb[sheet_name]
        header_row = find_header_row(ws)
        header_values = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
        headers = ["" if h is None else str(h).strip() for h in header_values]

        org_idx = get_col_idx(headers, ("organization", "organisation", "response organisation"))
        acct_idx = get_col_idx(headers, ("account number", "account no", "account"))
        rowname_idx = get_col_idx(headers, ("rowname",))

        answer_col_idxs = []
        for idx, col_name in enumerate(headers):
            if not col_name:
                continue
            low_col = normalize_text(col_name)
            if any(term in low_col for term in metadata_terms):
                continue
            answer_col_idxs.append(idx)

        if not answer_col_idxs:
            continue

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            org_val = row[org_idx] if (org_idx is not None and org_idx < len(row)) else None
            acct_val = row[acct_idx] if (acct_idx is not None and acct_idx < len(row)) else None
            row_name = row[rowname_idx] if (rowname_idx is not None and rowname_idx < len(row)) else None

            keep = False
            if acct_val is not None and account_ids and str(acct_val).strip() in account_ids:
                keep = True
            else:
                org_norm = normalize_text(org_val)
                if org_norm and (org_norm == org_norm_target or org_norm_target in org_norm or org_norm in org_norm_target):
                    keep = True
                elif matched_org_labels and org_val in matched_org_labels:
                    keep = True

            if not keep:
                continue

            for idx in answer_col_idxs:
                if idx >= len(row):
                    continue
                answer_val = row[idx]
                if is_missing(answer_val):
                    continue

                col_header = headers[idx]
                q_code = extract_question_code(sheet_name=sheet_name, column_header=col_header)
                q_key = canonical_question_key(q_code)
                q_text = parse_legacy_column_text(col_header)
                theme_tags = categorize_theme(sheet_name=sheet_name, question_text=q_text, answer=answer_val)

                records.append(
                    {
                        "year": year,
                        "source": "legacy_xlsx",
                        "source_file": str(file_path),
                        "org_name_matched": org_val,
                        "question_sheet": sheet_name,
                        "question_code": q_code,
                        "question_key": q_key,
                        "question_text": q_text,
                        "row_name": row_name,
                        "column_header": col_header,
                        "answer": answer_val,
                        "theme_tags": " | ".join(theme_tags),
                        "primary_theme": primary_theme(theme_tags),
                    }
                )

    return records, matched_org_labels


def extract_many_from_legacy_xlsx(
    year: int, file_path: Path, org_names: Sequence[str]
) -> Tuple[List[Dict[str, object]], Dict[str, List[str]]]:
    """Extract one CDP workbook for many organizations in a single scan."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    summary_sheet_name = (
        "Summary Data"
        if "Summary Data" in sheet_names
        else ("Summary" if "Summary" in sheet_names else None)
    )
    targets = [name for name in dict.fromkeys(org_names) if normalize_text(name)]
    matched_by_target: Dict[str, List[str]] = {name: [] for name in targets}
    target_by_account: Dict[str, str] = {}
    target_by_org_norm: Dict[str, str] = {}

    if summary_sheet_name:
        ws = wb[summary_sheet_name]
        header_row = find_header_row(ws)
        headers = [
            ws.cell(row=header_row, column=c).value
            for c in range(1, ws.max_column + 1)
        ]
        headers_str = ["" if value is None else str(value) for value in headers]
        org_idx = get_col_idx(
            headers_str,
            ("organization", "organisation", "response organisation"),
        )
        acct_idx = get_col_idx(
            headers_str, ("account number", "account no", "account")
        )
        if org_idx is not None:
            org_values = []
            acct_values = []
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                org_values.append(row[org_idx] if org_idx < len(row) else None)
                acct_values.append(
                    row[acct_idx]
                    if acct_idx is not None and acct_idx < len(row)
                    else None
                )
            for target in targets:
                matched_idx, labels = match_org_rows(org_values, target)
                matched_by_target[target] = labels
                for index in matched_idx:
                    org_norm = normalize_text(org_values[index])
                    if org_norm:
                        target_by_org_norm.setdefault(org_norm, target)
                    account = acct_values[index]
                    if account is not None and str(account).strip():
                        target_by_account.setdefault(str(account).strip(), target)

    records: List[Dict[str, object]] = []
    metadata_terms = (
        "account number",
        "account no",
        "organisation",
        "organization",
        "country",
        "access",
        "public",
        "samples",
        "response received date",
        "activities",
        "sectors",
        "industries",
        "primary activity",
        "primary sector",
        "primary industry",
        "primary questionnaire sector",
        "ticker",
        "isin",
        "row",
        "rowname",
        "region",
        "status",
        "submitted date",
        "invitation status",
        "request response status",
        "attachments",
        "complexity",
        "version",
        "back to index",
    )
    for sheet_name in sheet_names:
        low_sheet = sheet_name.lower()
        if low_sheet in {"criteria", "response language"}:
            continue
        if low_sheet.startswith("summary"):
            continue

        ws = wb[sheet_name]
        header_row = find_header_row(ws)
        header_values = [
            ws.cell(row=header_row, column=c).value
            for c in range(1, ws.max_column + 1)
        ]
        headers = [
            "" if value is None else str(value).strip()
            for value in header_values
        ]
        org_idx = get_col_idx(
            headers, ("organization", "organisation", "response organisation")
        )
        acct_idx = get_col_idx(
            headers, ("account number", "account no", "account")
        )
        rowname_idx = get_col_idx(headers, ("rowname",))
        answer_col_idxs = [
            index
            for index, column_name in enumerate(headers)
            if column_name
            and not any(
                term in normalize_text(column_name) for term in metadata_terms
            )
        ]
        if not answer_col_idxs:
            continue

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            org_value = (
                row[org_idx]
                if org_idx is not None and org_idx < len(row)
                else None
            )
            account_value = (
                row[acct_idx]
                if acct_idx is not None and acct_idx < len(row)
                else None
            )
            account_key = (
                str(account_value).strip() if account_value is not None else ""
            )
            requested_org = target_by_account.get(account_key)
            if requested_org is None:
                requested_org = target_by_org_norm.get(normalize_text(org_value))
            if requested_org is None:
                continue
            row_name = (
                row[rowname_idx]
                if rowname_idx is not None and rowname_idx < len(row)
                else None
            )
            for index in answer_col_idxs:
                if index >= len(row) or is_missing(row[index]):
                    continue
                answer = row[index]
                column_header = headers[index]
                question_code = extract_question_code(
                    sheet_name=sheet_name, column_header=column_header
                )
                question_text = parse_legacy_column_text(column_header)
                theme_tags = categorize_theme(
                    sheet_name=sheet_name,
                    question_text=question_text,
                    answer=answer,
                )
                records.append(
                    {
                        "year": year,
                        "source": "legacy_xlsx",
                        "source_file": str(file_path),
                        "requested_org_name": requested_org,
                        "org_name_matched": org_value,
                        "question_sheet": sheet_name,
                        "question_code": question_code,
                        "question_key": canonical_question_key(question_code),
                        "question_text": question_text,
                        "row_name": row_name,
                        "column_header": column_header,
                        "answer": answer,
                        "theme_tags": " | ".join(theme_tags),
                        "primary_theme": primary_theme(theme_tags),
                    }
                )
    wb.close()
    return records, matched_by_target


def extract_from_2024_parquet(
    year: int, file_path: Path, org_name: str
) -> Tuple[List[Dict[str, object]], List[str]]:
    df = pd.read_parquet(file_path)
    if "disclosing_organization" not in df.columns:
        raise ValueError(f"'disclosing_organization' column missing in {file_path}")

    df = df.copy()
    df["org_norm"] = df["disclosing_organization"].map(normalize_text)
    target = normalize_text(org_name)

    mask = (
        (df["org_norm"] == target)
        | (df["org_norm"].str.contains(re.escape(target), na=False))
        | df["org_norm"].map(lambda x: bool(x) and target in x)
    )
    matched = df[mask].copy()

    matched_labels: List[str] = sorted(matched["disclosing_organization"].dropna().astype(str).unique().tolist())
    if matched.empty:
        unique_org = sorted(df["org_norm"].dropna().unique().tolist())
        nearest = get_close_matches(target, unique_org, n=5, cutoff=0.86)
        if nearest:
            matched = df[df["org_norm"].isin(nearest)].copy()
            matched_labels = sorted(matched["disclosing_organization"].dropna().astype(str).unique().tolist())

    if matched.empty:
        return [], []

    out: List[Dict[str, object]] = []
    use_cols = [c for c in ["question_number", "question_text", "row_name", "column_header", "content_full"] if c in matched.columns]
    for row in matched[["disclosing_organization"] + use_cols].itertuples(index=False):
        row_dict = row._asdict()
        answer = row_dict.get("content_full")
        if is_missing(answer):
            continue

        q_number = str(row_dict.get("question_number", "") or "")
        q_text = str(row_dict.get("question_text", "") or "")
        q_code = extract_question_code(sheet_name=q_number, column_header=q_number or q_text)
        q_key = canonical_question_key(q_code)
        theme_tags = categorize_theme(sheet_name=q_number, question_text=q_text, answer=answer)
        out.append(
            {
                "year": year,
                "source": "responses_parquet",
                "source_file": str(file_path),
                "org_name_matched": row_dict.get("disclosing_organization"),
                "question_sheet": q_number,
                "question_code": q_code,
                "question_key": q_key,
                "question_text": q_text,
                "row_name": row_dict.get("row_name"),
                "column_header": row_dict.get("column_header"),
                "answer": answer,
                "theme_tags": " | ".join(theme_tags),
                "primary_theme": primary_theme(theme_tags),
            }
        )
    return out, matched_labels


def extract_many_from_2024_parquet(
    year: int, file_path: Path, org_names: Sequence[str]
) -> Tuple[List[Dict[str, object]], Dict[str, List[str]]]:
    """Extract one CDP response parquet for many organizations."""
    df = pd.read_parquet(file_path)
    if "disclosing_organization" not in df.columns:
        raise ValueError(
            f"'disclosing_organization' column missing in {file_path}"
        )
    df = df.copy()
    df["org_norm"] = df["disclosing_organization"].map(normalize_text)
    unique_orgs = sorted(df["org_norm"].dropna().unique().tolist())
    matched_parts = []
    matched_by_target: Dict[str, List[str]] = {}
    for org_name in dict.fromkeys(org_names):
        target = normalize_text(org_name)
        mask = (
            df["org_norm"].eq(target)
            | df["org_norm"].str.contains(re.escape(target), na=False)
            | df["org_norm"].map(lambda value: bool(value) and value in target)
        )
        matched = df[mask].copy()
        if matched.empty:
            nearest = get_close_matches(target, unique_orgs, n=1, cutoff=0.86)
            if nearest:
                matched = df[df["org_norm"].isin(nearest)].copy()
        labels = sorted(
            matched["disclosing_organization"]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
        )
        matched_by_target[org_name] = labels
        if not matched.empty:
            matched["requested_org_name"] = org_name
            matched_parts.append(matched)

    if not matched_parts:
        return [], matched_by_target
    matched = pd.concat(matched_parts, ignore_index=True)
    use_cols = [
        column
        for column in [
            "question_number",
            "question_text",
            "row_name",
            "column_header",
            "content_full",
        ]
        if column in matched.columns
    ]
    records: List[Dict[str, object]] = []
    for row in matched[
        ["requested_org_name", "disclosing_organization", *use_cols]
    ].itertuples(index=False):
        row_dict = row._asdict()
        answer = row_dict.get("content_full")
        if is_missing(answer):
            continue
        question_number = str(row_dict.get("question_number", "") or "")
        question_text = str(row_dict.get("question_text", "") or "")
        question_code = extract_question_code(
            sheet_name=question_number,
            column_header=question_number or question_text,
        )
        theme_tags = categorize_theme(
            sheet_name=question_number,
            question_text=question_text,
            answer=answer,
        )
        records.append(
            {
                "year": year,
                "source": "responses_parquet",
                "source_file": str(file_path),
                "requested_org_name": row_dict.get("requested_org_name"),
                "org_name_matched": row_dict.get("disclosing_organization"),
                "question_sheet": question_number,
                "question_code": question_code,
                "question_key": canonical_question_key(question_code),
                "question_text": question_text,
                "row_name": row_dict.get("row_name"),
                "column_header": row_dict.get("column_header"),
                "answer": answer,
                "theme_tags": " | ".join(theme_tags),
                "primary_theme": primary_theme(theme_tags),
            }
        )
    return records, matched_by_target


def build_question_matched_table(df_long: pd.DataFrame, start_year: int, end_year: int) -> pd.DataFrame:
    if df_long.empty:
        return df_long.copy()

    grouped = (
        df_long.assign(answer_str=df_long["answer"].astype(str))
        .groupby(["question_key", "question_code", "question_text", "year"], dropna=False)["answer_str"]
        .apply(lambda s: " || ".join(sorted(set(v for v in s if v and v.lower() != "nan"))[:5]))
        .reset_index(name="answers")
    )
    pivot = grouped.pivot_table(
        index=["question_key", "question_code", "question_text"],
        columns="year",
        values="answers",
        aggfunc="first",
    ).reset_index()

    for year in range(start_year, end_year + 1):
        if year not in pivot.columns:
            pivot[year] = ""

    year_cols = [year for year in range(start_year, end_year + 1)]
    pivot = pivot[["question_key", "question_code", "question_text"] + year_cols]
    pivot = pivot.sort_values(["question_key", "question_code", "question_text"], na_position="last")
    pivot.columns = [f"year_{c}" if isinstance(c, int) else c for c in pivot.columns]
    return pivot


def build_theme_summary(df_long: pd.DataFrame, start_year: int, end_year: int) -> pd.DataFrame:
    if df_long.empty:
        return df_long.copy()

    summary = (
        df_long.assign(primary_theme=df_long["primary_theme"].fillna("other"))
        .groupby(["primary_theme", "year"], dropna=False)
        .size()
        .reset_index(name="answer_count")
    )
    pivot = summary.pivot_table(index="primary_theme", columns="year", values="answer_count", aggfunc="sum", fill_value=0).reset_index()
    for year in range(start_year, end_year + 1):
        if year not in pivot.columns:
            pivot[year] = 0
    pivot = pivot[["primary_theme"] + [year for year in range(start_year, end_year + 1)]]
    pivot.columns = ["primary_theme"] + [f"year_{year}" for year in range(start_year, end_year + 1)]
    return pivot.sort_values("primary_theme")


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract CDP questionnaire answers across years (2014-2024).")
    parser.add_argument("--org-name", help="Organization name to search for")
    parser.add_argument("--factset-id", help="FactSet entity ID to resolve to an organization name before CDP extraction")
    parser.add_argument("--cdp-root", default=str(DEFAULT_CDP_ROOT), help="Root folder for CDP yearly data")
    parser.add_argument("--start-year", type=int, default=2014)
    parser.add_argument("--end-year", type=int, default=2024)
    parser.add_argument("--out-dir", default=str(DEFAULT_OUT_DIR))
    args = parser.parse_args()

    if not args.org_name and not args.factset_id:
        parser.error("Either --org-name or --factset-id must be provided")

    org_name = args.org_name
    if not org_name:
        org_name = resolve_org_name_from_factset_id(args.factset_id)

    cdp_root = Path(args.cdp_root)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if not cdp_root.exists():
        raise SystemExit(f"CDP root path not found: {cdp_root}")
    if args.start_year > args.end_year:
        raise SystemExit("start-year must be <= end-year")

    all_records: List[Dict[str, object]] = []
    matched_orgs_by_year: Dict[int, List[str]] = {}

    for year in range(args.start_year, args.end_year + 1):
        year_dir = cdp_root / str(year)
        file_path = parse_year_file(year_dir)
        if file_path is None:
            LOGGER.warning("Skipping %s (no supported CDP file found)", year)
            continue

        LOGGER.info("Processing %s from %s", year, file_path.name)
        try:
            if file_path.suffix.lower() == ".parquet":
                records, matched_orgs = extract_from_2024_parquet(year, file_path, org_name)
            else:
                records, matched_orgs = extract_from_legacy_xlsx(year, file_path, org_name)
        except Exception as exc:
            LOGGER.error("Failed year %s (%s): %s", year, file_path, exc)
            continue

        all_records.extend(records)
        matched_orgs_by_year[year] = matched_orgs
        LOGGER.info("Year %s: matched_orgs=%d, extracted_answers=%d", year, len(matched_orgs), len(records))

    if not all_records:
        LOGGER.warning("No answers found for organization '%s' in %s-%s", org_name, args.start_year, args.end_year)
        return

    df_long = pd.DataFrame(all_records)
    df_long = df_long.sort_values(["year", "question_key", "question_code"], na_position="last")

    matched_table = build_question_matched_table(df_long, args.start_year, args.end_year)
    theme_summary = build_theme_summary(df_long, args.start_year, args.end_year)

    slug = slugify(org_name)
    long_out = out_dir / f"cdp_org_answers_{slug}_{args.start_year}_{args.end_year}_long.csv"
    matched_out = out_dir / f"cdp_org_answers_{slug}_{args.start_year}_{args.end_year}_matched_questions.csv"
    orgs_out = out_dir / f"cdp_org_answers_{slug}_{args.start_year}_{args.end_year}_matched_orgs_by_year.csv"
    theme_out = out_dir / f"cdp_org_answers_{slug}_{args.start_year}_{args.end_year}_themes.csv"

    df_long.to_csv(long_out, index=False)
    matched_table.to_csv(matched_out, index=False)
    theme_summary.to_csv(theme_out, index=False)

    org_rows = []
    for y in range(args.start_year, args.end_year + 1):
        labels = matched_orgs_by_year.get(y, [])
        org_rows.append({"year": y, "matched_organizations": " || ".join(labels)})
    pd.DataFrame(org_rows).to_csv(orgs_out, index=False)

    LOGGER.info("Wrote long answers: %s (rows=%d)", long_out, len(df_long))
    LOGGER.info("Wrote matched question table: %s (rows=%d)", matched_out, len(matched_table))
    LOGGER.info("Wrote theme summary: %s", theme_out)
    LOGGER.info("Wrote matched organization list: %s", orgs_out)


if __name__ == "__main__":
    main()
