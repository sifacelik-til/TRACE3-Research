"""Build a tier-3 upstream/downstream supply-chain network for one target company.

Outputs (per target):
- network_<target>/nodes.csv
- network_<target>/edges.csv
- network_<target>/trucost_glossary.csv

The network keeps weighted edges (revenue_pct) and walks upstream/downstream to tier 3.
It also enriches nodes with FactSet industry sector buckets from the image mapping and
Trucost Scope 1/2/3 fields using the glossary workbook.
"""
from __future__ import annotations

from argparse import ArgumentParser
from pathlib import Path
import logging
import re

import numpy as np
import openpyxl
import pandas as pd

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
LOGGER = logging.getLogger(__name__)

DATA_RAW = Path("data") / "raw"
FACTSET_DIR = DATA_RAW / "FactSet"
TRUCOST_DIR = DATA_RAW / "Trucost (Access through WRDS)"
OUT_DIR = Path("data") / "processed"
OUT_DIR.mkdir(parents=True, exist_ok=True)

TRUCOST_GLOSSARY = TRUCOST_DIR / "trucost_environmental_data_item_list.xlsx"
TRUCOST_CSV_CANDIDATE = TRUCOST_DIR / "260710 trucost pulic-ghg-2011to24.csv"

RELATIONSHIP_FILE = FACTSET_DIR / "ent_supply_chain_v1_full_3856" / "ent_scr_relationships.txt"
SYM_ENTITY_FILE = FACTSET_DIR / "sym_entity_v1_full_12328" / "sym_entity.txt"
SYM_SECTOR_FILE = FACTSET_DIR / "sym_entity_v1_full_12328" / "sym_entity_sector.txt"
INDUSTRY_MAP_FILE = FACTSET_DIR / "ref_hub_v2_full_3565" / "factset_industry_map.txt"

SELECTED_TRUCOST_IDS = [
    319410, 319411, 319413, 376883, 319414, 376884, 367750, 376886, 378459,
    368653, 377767, 368654, 377768, 368655, 377405, 368704, 377406, 368705,
    377407, 368706, 377408, 378680, 378681, 378682, 378683, 378684, 378685,
    368652, 377766, 326737, 368707, 377409, 378686, 319415, 319412,
]
SELECTED_TRUCOST_COLS = [f"di_{i}" for i in SELECTED_TRUCOST_IDS]

FACTSET_SECTOR_BUCKETS = [
    ("Non-Energy Minerals", {"1105", "1115", "1120", "1125", "1130", "1135"}),
    ("Producer Manufacturing", {"1205", "1210", "1220", "1225", "1230", "1235", "1245", "1250", "1255"}),
    ("Electronic Technology", {"1305", "1310", "1315", "1320", "1330", "1340", "1345", "1352", "1355"}),
    ("Consumer Durables", {"1405", "1410", "1415", "1420", "1425", "1430", "1435", "1445"}),
    ("Energy Minerals", {"2105", "2110", "2120", "2125"}),
    ("Process Industries", {"2205", "2210", "2215", "2220", "2225", "2230", "2235", "2240"}),
    ("Health Technology", {"2305", "2310", "2315", "2320", "2325"}),
    ("Consumer Non-Durables", {"2405", "2410", "2415", "2420", "2425", "2430", "2435", "2440", "2450"}),
    ("Industrial Services", {"3105", "3110", "3115", "3120", "3130"}),
    ("Commercial Services", {"3205", "3210", "3215", "3235"}),
    ("Distribution Services", {"3255", "3260", "3265", "3270"}),
    ("Technology Services", {"3305", "3308", "3310", "3320"}),
    ("Health Services", {"3355", "3360", "3365", "3370"}),
    ("Consumer Services", {"3405", "3410", "3415", "3420", "3425", "3430", "3435", "3440", "3445", "3450"}),
    ("Retail Trade", {"3505", "3510", "3515", "3520", "3525", "3530", "3535", "3540", "3545", "3550"}),
    ("Transportation", {"4605", "4610", "4615", "4620", "4625", "4630"}),
    ("Utilities", {"4705", "4735", "4755", "4760"}),
    ("Finance", {"4805", "4810", "4825", "4830", "4840", "4845", "4850", "4855", "4860", "4865", "4875", "4880", "4885", "4890"}),
    ("Communications", {"4905", "4910", "4915"}),
    ("Miscellaneous", {"6005", "6010"}),
]


def read_pipe(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path, sep="|", quotechar='"', dtype=str, encoding="utf-8")
    except Exception:
        return pd.read_csv(path, sep="|", quotechar='"', dtype=str, encoding="latin-1")


def normalize_cols(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [c.strip().strip('"') for c in df.columns]
    return df


def slugify(text: str) -> str:
    text = (text or "").lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_")


def sector_bucket(industry_code: str, industry_desc: str) -> str:
    code = str(industry_code or "").strip()
    for bucket, codes in FACTSET_SECTOR_BUCKETS:
        if code in codes:
            return bucket
    desc = (industry_desc or "").lower()
    if any(k in desc for k in ("semiconductor", "electronics", "computer", "telecommunications", "aerospace", "defense", "technology")):
        return "Electronic Technology"
    return "Other"


def load_trucost_glossary(path: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        try:
            did = int(float(row[3])) if row[3] is not None and str(row[3]).strip() else None
        except Exception:
            did = None
        if did in SELECTED_TRUCOST_IDS:
            rows.append({
                "dataItemId": did,
                "dataItemName": row[2],
                "definition": row[4],
            })
    wb.close()
    return pd.DataFrame(rows).sort_values("dataItemId")


def parse_args():
    parser = ArgumentParser()
    parser.add_argument("--target", required=True, help="Target company name or FactSet entity id")
    parser.add_argument("--max-tier", type=int, default=3, choices=[1, 2, 3], help="Max tier depth")
    parser.add_argument("--output-dir", default=None, help="Optional output directory")
    return parser.parse_args()


def select_target_entity(sym: pd.DataFrame, target: str, rel: pd.DataFrame | None = None) -> pd.Series:
    target_norm = target.strip().lower()
    exact_id = sym[sym["FACTSET_ENTITY_ID"] == target.strip()]
    if not exact_id.empty:
        return exact_id.iloc[0]
    exact_name = sym[sym["ENTITY_PROPER_NAME"].fillna("").astype(str).str.strip().str.lower() == target_norm]
    if not exact_name.empty:
        return exact_name.iloc[0]
    contains = sym[sym["ENTITY_PROPER_NAME_NORM"].str.contains(re.escape(target_norm), na=False)].copy()
    if not contains.empty:
        contains["rank_startswith"] = contains["ENTITY_PROPER_NAME_NORM"].str.startswith(target_norm).astype(int)
        contains["rank_score"] = contains["ENTITY_PROPER_NAME_NORM"].astype(str).apply(lambda s: __import__("rapidfuzz").fuzz.token_sort_ratio(target_norm, s))
        contains["rank_len"] = contains["ENTITY_PROPER_NAME_NORM"].astype(str).str.len()
        if rel is not None and not rel.empty:
            candidate_ids = contains["FACTSET_ENTITY_ID"].astype(str).tolist()
            rel_counts = (
                pd.concat([
                    rel[rel["SOURCE_ID"].isin(candidate_ids)]["SOURCE_ID"],
                    rel[rel["TARGET_ID"].isin(candidate_ids)]["TARGET_ID"],
                ])
                .value_counts()
                .to_dict()
            )
            contains["rank_rel_count"] = contains["FACTSET_ENTITY_ID"].astype(str).map(rel_counts).fillna(0)
        else:
            contains["rank_rel_count"] = 0
        contains = contains.sort_values(["rank_rel_count", "rank_startswith", "rank_score", "rank_len"], ascending=[False, False, False, True])
        return contains.iloc[0]
    # lightweight fuzzy fallback
    from rapidfuzz import process, fuzz

    best = process.extractOne(target_norm, sym["ENTITY_PROPER_NAME_NORM"].dropna().astype(str).unique().tolist(), scorer=fuzz.token_sort_ratio)
    if best and best[1] >= 80:
        hit = sym[sym["ENTITY_PROPER_NAME_NORM"] == best[0]]
        if not hit.empty:
            return hit.iloc[0]
    raise SystemExit(f"Could not find target company: {target}")


def relation_neighbors(row, current_id: str, direction: str):
    rel_type = str(row.REL_TYPE).upper()
    source = str(row.SOURCE_ID)
    target = str(row.TARGET_ID)
    revenue_pct = pd.to_numeric(row.REVENUE_PCT_NUM, errors="coerce")
    start_date = row.START_DATE_DT
    end_date = row.END_DATE_DT

    # Interpret relationship from the current node's perspective.
    if current_id == source:
        if direction == "downstream" and "CUSTOMER" in rel_type:
            return target, source, target
        if direction == "upstream" and "SUPPLIER" in rel_type:
            return target, target, source
    if current_id == target:
        if direction == "upstream" and "CUSTOMER" in rel_type:
            return source, source, target
        if direction == "downstream" and "SUPPLIER" in rel_type:
            return source, target, source
    return None


def get_company_row(sym: pd.DataFrame, company_id: str) -> pd.Series | None:
    hit = sym[sym["FACTSET_ENTITY_ID"] == company_id]
    if hit.empty:
        return None
    return hit.iloc[0]


def main():
    args = parse_args()
    run_out_dir = Path(args.output_dir) if args.output_dir else (OUT_DIR / f"network_{slugify(args.target)}")

    LOGGER.info("Loading FactSet entities")
    sym = normalize_cols(read_pipe(SYM_ENTITY_FILE))
    sym = sym.rename(columns={sym.columns[0]: "FACTSET_ENTITY_ID", sym.columns[1]: "ENTITY_PROPER_NAME"})
    sym["ENTITY_PROPER_NAME_NORM"] = sym["ENTITY_PROPER_NAME"].fillna("").astype(str).str.strip().str.lower()
    sym = sym.drop_duplicates(subset=["FACTSET_ENTITY_ID"], keep="first").copy()

    LOGGER.info("Loading industry metadata")
    sector = normalize_cols(read_pipe(SYM_SECTOR_FILE))
    sector = sector.rename(columns={sector.columns[0]: "FACTSET_ENTITY_ID", sector.columns[2]: "INDUSTRY_CODE", sector.columns[3]: "SECTOR_CODE"})
    sector = sector[["FACTSET_ENTITY_ID", "INDUSTRY_CODE"]].drop_duplicates("FACTSET_ENTITY_ID", keep="first")
    ind_map = normalize_cols(read_pipe(INDUSTRY_MAP_FILE))
    ind_map = ind_map.rename(columns={ind_map.columns[0]: "INDUSTRY_CODE", ind_map.columns[1]: "INDUSTRY_DESC"})

    companies = sym[["FACTSET_ENTITY_ID", "ENTITY_PROPER_NAME", "ENTITY_PROPER_NAME_NORM"]].merge(
        sector, on="FACTSET_ENTITY_ID", how="left"
    ).merge(ind_map[["INDUSTRY_CODE", "INDUSTRY_DESC"]], on="INDUSTRY_CODE", how="left")
    companies["INDUSTRY_CODE"] = companies["INDUSTRY_CODE"].fillna("9999")
    companies["INDUSTRY_DESC"] = companies["INDUSTRY_DESC"].fillna("Not Classified")
    companies = companies[companies["INDUSTRY_CODE"].astype(str) != "9999"].copy()
    companies["sector_bucket"] = companies.apply(lambda r: sector_bucket(r["INDUSTRY_CODE"], r["INDUSTRY_DESC"]), axis=1)

    LOGGER.info("Loading relationships")
    rel = normalize_cols(read_pipe(RELATIONSHIP_FILE))
    rel = rel.rename(columns={rel.columns[0]: "ID", rel.columns[1]: "REL_TYPE", rel.columns[2]: "SOURCE_FACTSET_ENTITY_ID", rel.columns[3]: "TARGET_FACTSET_ENTITY_ID"})
    for c in ("START_DATE", "END_DATE", "REVENUE_PCT"):
        if c not in rel.columns:
            rel[c] = None
    rel["SOURCE_ID"] = rel["SOURCE_FACTSET_ENTITY_ID"].fillna("").astype(str).str.strip().str.strip('"')
    rel["TARGET_ID"] = rel["TARGET_FACTSET_ENTITY_ID"].fillna("").astype(str).str.strip().str.strip('"')
    rel["REL_TYPE_UP"] = rel["REL_TYPE"].fillna("").astype(str).str.upper()
    rel["REVENUE_PCT_NUM"] = pd.to_numeric(rel["REVENUE_PCT"], errors="coerce")
    rel["START_DATE_DT"] = pd.to_datetime(rel["START_DATE"], errors="coerce")
    rel["END_DATE_DT"] = pd.to_datetime(rel["END_DATE"], errors="coerce")

    target_row = select_target_entity(sym, args.target, rel)
    target_id = target_row["FACTSET_ENTITY_ID"]
    target_name = target_row["ENTITY_PROPER_NAME"]
    LOGGER.info("Target selected: %s (%s)", target_name, target_id)

    # BFS for upstream/downstream tiers
    nodes = {
        target_id: {
            "node_id": target_id,
            "node_name": target_name,
            "tier_upstream": 0,
            "tier_downstream": 0,
            "is_target": True,
        }
    }
    edge_rows = []
    visited_up = {target_id}
    visited_down = {target_id}
    frontier_up = {target_id: 1.0}
    frontier_down = {target_id: 1.0}

    def add_node(node_id: str, direction: str, tier: int):
        rec = nodes.setdefault(node_id, {"node_id": node_id, "node_name": None, "tier_upstream": None, "tier_downstream": None, "is_target": False})
        key = f"tier_{direction}"
        current = rec.get(key)
        if current is None or tier < current:
            rec[key] = tier

    for direction in ("upstream", "downstream"):
        frontier = frontier_up if direction == "upstream" else frontier_down
        visited = visited_up if direction == "upstream" else visited_down
        for tier in range(1, args.max_tier + 1):
            frontier_ids = list(frontier.keys())
            if not frontier_ids:
                break
            mask = rel["SOURCE_ID"].isin(frontier_ids) | rel["TARGET_ID"].isin(frontier_ids)
            sub = rel[mask]
            next_frontier = {}
            next_weights = {}
            for row in sub.itertuples(index=False):
                for current_id, current_weight in frontier.items():
                    neighbor_info = relation_neighbors(row, current_id, direction)
                    if not neighbor_info:
                        continue
                    neighbor_id, edge_source, edge_target = neighbor_info
                    if not neighbor_id:
                        continue
                    edge_weight_pct = pd.to_numeric(row.REVENUE_PCT_NUM, errors="coerce")
                    if pd.notna(current_weight) and pd.notna(edge_weight_pct):
                        path_weight_pct = current_weight * (float(edge_weight_pct) / 100.0)
                    else:
                        path_weight_pct = np.nan
                    edge_rows.append({
                        "root_target_id": target_id,
                        "root_target_name": target_name,
                        "direction": direction,
                        "tier": tier,
                        "source_id": edge_source,
                        "target_id": edge_target,
                        "neighbor_id": neighbor_id,
                        "relation_type": row.REL_TYPE,
                        "revenue_pct": edge_weight_pct,
                        "path_weight_pct": path_weight_pct * 100.0 if pd.notna(path_weight_pct) else np.nan,
                        "start_date": row.START_DATE_DT,
                        "end_date": row.END_DATE_DT,
                    })
                    add_node(neighbor_id, direction, tier)
                    # keep the first path weight found for the next tier expansion
                    if neighbor_id not in visited:
                        visited.add(neighbor_id)
                        next_frontier[neighbor_id] = path_weight_pct if pd.notna(path_weight_pct) else current_weight
            frontier = next_frontier
            if direction == "upstream":
                frontier_up = frontier
            else:
                frontier_down = frontier

    # Build node table
    node_ids = list(nodes.keys())
    node_df = companies[companies["FACTSET_ENTITY_ID"].isin(node_ids)].copy()
    node_df = node_df.rename(columns={
        "FACTSET_ENTITY_ID": "node_id",
        "ENTITY_PROPER_NAME": "node_name",
        "INDUSTRY_CODE": "industry_code",
        "INDUSTRY_DESC": "industry_desc",
    })
    node_df["sector_bucket"] = node_df.apply(lambda r: sector_bucket(r["industry_code"], r["industry_desc"]), axis=1)
    node_df["is_target"] = node_df["node_id"].eq(target_id)
    node_df["tier_upstream"] = node_df["node_id"].map(lambda x: nodes.get(x, {}).get("tier_upstream"))
    node_df["tier_downstream"] = node_df["node_id"].map(lambda x: nodes.get(x, {}).get("tier_downstream"))

    if target_id not in node_df["node_id"].tolist():
        target_row = pd.DataFrame([
            {
                "node_id": target_id,
                "node_name": target_name,
                "industry_code": None,
                "industry_desc": None,
                "sector_bucket": "Other",
                "is_target": True,
                "tier_upstream": 0,
                "tier_downstream": 0,
            }
        ])
        node_df = pd.concat([node_df, target_row], ignore_index=True)
        node_df["is_target"] = node_df["node_id"].eq(target_id)

    # CDP exact matching
    cdp_summary = pd.DataFrame()
    cdp_id_col = None
    cdp_candidates = list(DATA_RAW.rglob("*full_extract*_isin*_summary*.parquet"))
    if not cdp_candidates:
        cdp_candidates = list(DATA_RAW.rglob("*full_extract*summary*.parquet"))
    if cdp_candidates:
        cdp_summary = pd.read_parquet(cdp_candidates[0])
        for c in ("disclosing_organization", "disclosing_organization_name", "disclosing_organization_full", "organization_name"):
            if c in cdp_summary.columns:
                cdp_summary["cdp_name_norm"] = cdp_summary[c].fillna("").astype(str).str.strip().str.lower()
                break
        if "isin" in cdp_summary.columns:
            cdp_summary["isin"] = cdp_summary["isin"].fillna("").astype(str).str.strip().str.upper()
        cdp_candidates_cols = [c for c in cdp_summary.columns if any(k in c.lower() for k in ("organization_id", "participant", "cdp_id", "company_id", "disclosing_organization_id"))]
        if cdp_candidates_cols:
            cdp_id_col = cdp_candidates_cols[0]
    isin_cols = [c for c in sym.columns if "isin" in c.lower()]
    ticker_cols = [c for c in sym.columns if c.lower() in ("ticker", "symbol", "primary_symbol") or "ticker" in c.lower() or "symbol" in c.lower()]
    meta = sym[["FACTSET_ENTITY_ID"] + isin_cols + ticker_cols].drop_duplicates("FACTSET_ENTITY_ID", keep="first").rename(columns={"FACTSET_ENTITY_ID": "node_id"})
    node_df = node_df.merge(meta, on="node_id", how="left")

    if not cdp_summary.empty and cdp_id_col is not None:
        if "isin" in cdp_summary.columns and isin_cols:
            isin_map = cdp_summary[["isin", cdp_id_col]].dropna(subset=["isin"]).drop_duplicates("isin", keep="first")
            isin_map = isin_map.rename(columns={"isin": "node_isin", cdp_id_col: "cdp_id"})
            node_df["node_isin"] = node_df[isin_cols].bfill(axis=1).iloc[:, 0] if isin_cols else pd.NA
            node_df = node_df.merge(isin_map, on="node_isin", how="left")
        if "cdp_name_norm" in cdp_summary.columns:
            name_map = cdp_summary[["cdp_name_norm", cdp_id_col]].dropna(subset=["cdp_name_norm"]).drop_duplicates("cdp_name_norm", keep="first")
            name_map = name_map.rename(columns={"cdp_name_norm": "node_name_norm", cdp_id_col: "cdp_id_name"})
            node_df["node_name_norm"] = node_df["node_name"].fillna("").astype(str).str.strip().str.lower()
            node_df = node_df.merge(name_map, on="node_name_norm", how="left")
            node_df["cdp_id"] = node_df.get("cdp_id").combine_first(node_df.get("cdp_id_name"))

    # Trucost exact matching and selected yearly fields
    tr = pd.DataFrame()
    glossary = load_trucost_glossary(TRUCOST_GLOSSARY)
    if TRUCOST_CSV_CANDIDATE.exists():
        tr = pd.read_csv(TRUCOST_CSV_CANDIDATE, dtype=str, encoding="latin-1")
        if "ticker" in tr.columns:
            tr["ticker"] = tr["ticker"].fillna("").astype(str).str.strip().str.upper()
        if "companyname" in tr.columns:
            tr["companyname_norm"] = tr["companyname"].fillna("").astype(str).str.strip().str.lower()
        if "companyid" in tr.columns:
            tr["companyid"] = tr["companyid"].fillna("").astype(str).str.strip()
        if "fiscalyear" not in tr.columns and "periodenddate" in tr.columns:
            tr["fiscalyear"] = pd.to_datetime(tr["periodenddate"], errors="coerce").dt.year

        tr_match = node_df[["node_id", "node_name"]].copy()
        tr_match = tr_match.merge(sym[["FACTSET_ENTITY_ID"] + ticker_cols].rename(columns={"FACTSET_ENTITY_ID": "node_id"}), on="node_id", how="left")
        tr_match["node_ticker"] = pd.NA
        for c in ticker_cols:
            tr_match["node_ticker"] = tr_match["node_ticker"].combine_first(tr_match[c].fillna("").astype(str).str.strip().str.upper())
        tr_match["node_name_norm"] = tr_match["node_name"].fillna("").astype(str).str.strip().str.lower()

        tr_ticker = pd.DataFrame()
        if "ticker" in tr.columns:
            tr_ticker = tr[["ticker", "companyid"]].dropna(subset=["ticker"]).drop_duplicates("ticker", keep="first")
            tr_ticker = tr_ticker.rename(columns={"ticker": "node_ticker", "companyid": "trucost_companyid"})
        tr_name = pd.DataFrame()
        if "companyname_norm" in tr.columns:
            tr_name = tr[["companyname_norm", "companyid"]].dropna(subset=["companyname_norm"]).drop_duplicates("companyname_norm", keep="first")
            tr_name = tr_name.rename(columns={"companyname_norm": "node_name_norm", "companyid": "trucost_companyid_name"})

        tr_match = tr_match.merge(tr_ticker, on="node_ticker", how="left") if not tr_ticker.empty else tr_match
        tr_match = tr_match.merge(tr_name, on="node_name_norm", how="left") if not tr_name.empty else tr_match
        if "trucost_companyid_name" in tr_match.columns:
            tr_match["trucost_companyid"] = tr_match["trucost_companyid"].combine_first(tr_match["trucost_companyid_name"])
        tr_match = tr_match[["node_id", "trucost_companyid"]].drop_duplicates("node_id", keep="first")
        node_df = node_df.merge(tr_match, on="node_id", how="left")

        available_di_cols = [c for c in SELECTED_TRUCOST_COLS if c in tr.columns]
        if available_di_cols and "fiscalyear" in tr.columns:
            tr["fiscalyear"] = pd.to_numeric(tr["fiscalyear"], errors="coerce")
            for c in available_di_cols:
                tr[c] = pd.to_numeric(tr[c], errors="coerce")
            
            tr_yearly = tr[
                tr["companyid"].isin(
                    node_df["trucost_companyid"].dropna().astype(str)
                )
            ].copy()

            if not tr_yearly.empty:
                tr_yearly = tr_yearly[
                    ["companyid", "fiscalyear"] + available_di_cols
                ].rename(
                    columns={
                        "companyid": "trucost_companyid",
                        "fiscalyear": "emissions_year"
                    }
                )

                node_df = node_df.merge(
                    tr_yearly,
                    on="trucost_companyid",
                    how="left"
                )


    # Scope summary columns from the requested Trucost set
    for src_col, out_col in {
        "di_319413": "scope1_gross",
        "di_376883": "scope1_as_reported",
        "di_319414": "scope2_location_based",
        "di_376884": "scope2_location_based_as_reported",
        "di_367750": "scope2_market_based",
        "di_376886": "scope2_market_based_as_reported",
        "di_319415": "scope3_upstream_total",
        "di_326737": "scope3_downstream_total",
    }.items():
        if src_col in node_df.columns:
            node_df[out_col] = node_df[src_col]

    # Save outputs
    out_dir = run_out_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    node_df.to_csv(out_dir / "nodes.csv", index=False)
    edge_df = pd.DataFrame(edge_rows).drop_duplicates()
    edge_df.to_csv(out_dir / "edges.csv", index=False)
    glossary.to_csv(out_dir / "trucost_glossary.csv", index=False)
    LOGGER.info("Wrote %d nodes and %d edges to %s", len(node_df), len(edge_df), out_dir)


if __name__ == "__main__":
    main()
